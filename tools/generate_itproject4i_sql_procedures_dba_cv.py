from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_DBA_Procedimientos_Almacenados_SQL_CV.docx"


def write(paragraph, value, size=9.2, bold=False, color="26333F"):
    run = paragraph.add_run(value)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def section(doc, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    write(p, label.upper(), 10.5, True, "174B70")


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.7)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, value)


def role(doc, title, company, dates, place, points, tech, break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = break_before
    write(p, f"{title} | {company}", 10, True, "17365D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    write(p, f"{dates} | {place}", 8.7, color="596878")
    for point in points:
        bullet(doc, point)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write(p, "Tecnologías: ", 8.6, True, "596878")
    write(p, tech, 8.6, color="596878")


doc = Document()
page = doc.sections[0]
page.top_margin = Inches(0.52)
page.bottom_margin = Inches(0.52)
page.left_margin = Inches(0.65)
page.right_margin = Inches(0.65)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 16, True, "17365D")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "DBA / Especialista SQL | Procedimientos almacenados y optimización", 10.1, True, "174B70")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
write(p, "San José, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.6, color="596878")

section(doc, "Perfil profesional")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.04
write(p, "Especialista en SQL Server con más de 15 años de experiencia en desarrollo y administración de bases de datos empresariales. Ha diseñado y optimizado procedimientos almacenados, consultas T-SQL y objetos de base de datos para lógica de negocio, aplicaciones y reportes. Experiencia en análisis de rendimiento, índices, modelado relacional, migraciones, validación de datos y documentación técnica. Colabora con equipos de desarrollo y negocio para resolver problemas de datos y mejorar la mantenibilidad de soluciones SQL existentes.")

section(doc, "Competencias técnicas")
for label, value in [
    ("Desarrollo SQL", "SQL Server, T-SQL, procedimientos almacenados, consultas complejas, joins, vistas y objetos de base de datos."),
    ("Optimización", "Análisis de consultas, estrategias de índices, ajuste de rendimiento y mejora de procesos SQL."),
    ("Administración y calidad", "Modelado relacional, migraciones, integridad de datos, validaciones, manejo de errores y documentación."),
    ("Integración", "SSIS, ETL, SSRS, SSAS, Power BI, PostgreSQL, MySQL y Snowflake SQL."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    write(p, f"{label}: ", 9.1, True)
    write(p, value, 9.1)

section(doc, "Experiencia profesional")
role(doc, "Snowflake Developer / Data Engineer", "ServiceTitan", "Sep 2024 - Actualidad", "Remoto / Costa Rica", [
    "Migra lógica de reportes desde C# hacia Snowflake SQL y mantiene modelos dbt para procesamiento y análisis de datos.",
    "Optimiza cargas SQL y dbt; redujo en 20% el tiempo de procesamiento del data warehouse en una primera fase.",
], "Snowflake SQL, dbt, modelado de datos, Kimball")
role(doc, "Data Engineer", "SMASH Costa Rica", "May 2021 - Sep 2024", "San José, Costa Rica", [
    "Diseñó aplicaciones ETL y flujos de datos con SQL Server y SSIS según requisitos de clientes, reduciendo trabajo manual hasta en 40%.",
    "Automatizó análisis y controles de validación para detectar anomalías y mejorar la exactitud de datos en 30%.",
], "SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI")
role(doc, "SQL Developer", "Intertec International", "Feb 2019 - Abr 2021", "San José, Costa Rica", [
    "Desarrolló y optimizó procedimientos almacenados, consultas SQL y flujos de base de datos que soportaban lógica empresarial y análisis.",
    "Redujo los tiempos de consulta hasta en 50% mediante optimización SQL, índices y análisis de rendimiento.",
    "Mejoró la exactitud de datos en más de 25% con controles de validación y manejo de errores en procesos clave.",
], "SQL Server, T-SQL, SSIS, Salesforce, MySQL")
role(doc, "DBA / SQL Developer", "EL Tiempo", "Sep 2020 - Nov 2020", "Colombia", [
    "Lideró la migración de diez bases de datos, coordinando validación, integridad de datos y continuidad operativa.",
], "SQL Server, SSIS, Azure, SSRS", break_before=True)
role(doc, "DBA / SQL & BI Consultant", "Gold Data Networks", "Ene 2016 - Feb 2020", "Ciudad de Panamá, Panamá", [
    "Construyó bases de datos relacionales y estructuras de tablas desde cero para aplicaciones web y reportes.",
    "Desarrolló objetos SQL y procedimientos almacenados para mejorar el rendimiento y apoyar procesos de backend.",
], "SQL Server, SSIS, SSRS, PostgreSQL, Power BI")
role(doc, "Data Warehouse DBA", "BAC Credomatic", "Nov 2017 - Ene 2019", "San José, Costa Rica", [
    "Definió y optimizó tablas, índices y vistas para cargas analíticas; mantuvo pipelines ETL de múltiples fuentes.",
], "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure")

section(doc, "Experiencia adicional")
bullet(doc, "EducaTablet: administró y desarrolló bases SQL Server y capacitó a desarrolladores en buenas prácticas de programación de datos.")
bullet(doc, "VIGEOSOFT: modeló bases OLTP y documentó cambios estructurales mediante diccionarios de datos.")
bullet(doc, "ACH Cloud Services, Bosal y Optica Caroni: desarrollo SQL, administración de bases de datos y reportes empresariales.")

section(doc, "Formación e idiomas")
p = doc.add_paragraph()
write(p, "Analista de Sistemas, Informática - IUT Dr. Federico Rivero Palacio (2000-2004).")
p = doc.add_paragraph()
write(p, "Formación adicional: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI.")
p = doc.add_paragraph()
write(p, "Español: nativo | Inglés: dominio profesional.")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "DBA Especialista SQL Procedimientos Almacenados CV"
doc.save(OUTPUT)
print(OUTPUT)

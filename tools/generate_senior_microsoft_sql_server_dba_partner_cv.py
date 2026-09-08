from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from generate_fullstack_principal_sql_dba_cv import (
    OUTPUT_DIR,
    add_bullet,
    add_heading,
    add_role,
    add_run,
    configure_document,
    set_cell_borders,
    set_cell_margins,
    set_cell_shading,
    set_table_width,
)


DOCX_PATH = OUTPUT_DIR / "Pedro_Gutierrez_Senior_Microsoft_SQL_Server_DBA_CV.docx"


def add_header(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("PEDRO JAVIER GUTIERREZ ARMAS")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(0)
    add_run(
        subtitle,
        "Senior Microsoft SQL Server DBA | T-SQL Performance | Data Platforms",
        bold=True,
        size=10.5,
    )

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(5)
    add_run(
        contact,
        "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13",
        size=9,
    )


def add_skills_table(doc):
    rows = [
        (
            "SQL Server Administration",
            "Microsoft SQL Server, production database support, database design, maintenance, migrations, operational continuity",
        ),
        (
            "Performance Engineering",
            "Advanced T-SQL, complex queries, stored procedures, indexing strategies, query optimization, performance tuning, troubleshooting",
        ),
        (
            "Data Architecture and BI",
            "OLTP modeling, data warehousing, ETL/ELT, SSIS, SSAS, SSRS, Power BI, data lakes, Kimball practices",
        ),
        (
            "Automation and Cloud",
            "PowerShell, Python, Azure-based environments, Snowflake, dbt, Git-based collaboration, AI-assisted development",
        ),
        (
            "Data Operations",
            "Data integrity, validation checks, error handling, data security practices, technical documentation, stakeholder coordination",
        ),
    ]

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Area"
    header[1].text = "Tools and Practices"
    for cell in header:
        set_cell_shading(cell, "1F4E79")
        set_cell_borders(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(8.8)

    for area, details in rows:
        cells = table.add_row().cells
        cells[0].text = area
        cells[1].text = details
        for cell in cells:
            set_cell_borders(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
        cells[0].paragraphs[0].runs[0].font.bold = True
    set_table_width(table, [1.75, 5.65])


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_header(doc)

    add_heading(doc, "PROFESSIONAL SUMMARY")
    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(4)
    summary.add_run(
        "Senior Microsoft SQL Server DBA, SQL Developer, and Data Engineer with 15+ years of experience across database administration, database development, production support, performance optimization, data migrations, ETL, data warehousing, and reporting solutions. Strong hands-on background in advanced T-SQL, stored procedures, indexing, query tuning, relational and OLTP modeling, data integrity, database security practices, PowerShell, SSIS, and Azure-based environments. Proven ability to lead database migration work, troubleshoot performance and data issues, coordinate with cross-functional stakeholders, document technical environments, establish database practices, and train development teams."
    )

    add_heading(doc, "TECHNICAL SKILLS")
    add_skills_table(doc)

    add_heading(doc, "PROFESSIONAL EXPERIENCE")
    add_role(
        doc,
        "Snowflake Developer / Data Engineer",
        "ServiceTitan",
        "September 2024 - Present",
        "Remote / Costa Rica",
        [
            "Develop and optimize Snowflake SQL and dbt models supporting scalable reporting, analytics, and reusable data assets.",
            "Reduced data warehouse processing time by 20% during an initial performance optimization phase.",
            "Created a Kimball-based data warehouse modeling proof of concept to improve structure, query behavior, and processing performance.",
            "Support code quality and collaborative development using Cursor and an AI-assisted pull request review bot.",
        ],
        "Snowflake, Snowflake SQL, dbt, MetricFlow, Kimball, SQL, Cursor, AI-assisted PR review",
    )
    add_role(
        doc,
        "Data Engineer",
        "SMASH Costa Rica",
        "May 2021 - September 2024",
        "San Jose, Costa Rica",
        [
            "Built and maintained production data workflows using SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI, and Excel.",
            "Designed ETL applications aligned with client requirements, reducing manual data processing time by up to 40%.",
            "Developed automated data validation and anomaly-detection workflows that improved validation accuracy by 30%.",
            "Prepared, transformed, modeled, and validated datasets for reporting and operational decision-making.",
        ],
        "SQL Server, SSIS, PowerShell, Python, Snowflake, Power BI, Pandas, NumPy, Excel",
    )
    add_role(
        doc,
        "SQL Developer",
        "Intertec International",
        "February 2019 - April 2021",
        "San Jose, Costa Rica",
        [
            "Developed and optimized stored procedures, complex T-SQL queries, and database workflows supporting production business logic.",
            "Reduced query execution times by up to 50% through SQL optimization, indexing strategies, and performance tuning.",
            "Troubleshot data and database workflow issues while consolidating multiple source systems into analytics-ready datasets.",
            "Improved data accuracy by more than 25% through validation checks and error-handling mechanisms.",
        ],
        "SQL Server, T-SQL, SSIS, Salesforce, MySQL",
    )

    doc.add_page_break()
    add_role(
        doc,
        "DBA / SQL Developer",
        "EL Tiempo",
        "September 2020 - November 2020",
        "Colombia",
        [
            "Led planning and execution of migration activities for 10 SQL Server databases, delivering each phase on time and within scope.",
            "Coordinated migration risk identification, data validation, deployment activities, and continuity of operations with cross-functional teams.",
            "Supported SQL Server and Azure-based database and reporting environments during production transitions.",
        ],
        "SQL Server, SSIS, Azure, SSRS",
    )
    add_role(
        doc,
        "DBA / SQL and BI Consultant",
        "Gold Data Networks",
        "January 2016 - February 2020",
        "Panama City, Panama",
        [
            "Designed relational databases, schemas, tables, stored procedures, and supporting database objects for operational applications and reporting.",
            "Improved application performance through custom database logic and maintainable backend structures.",
            "Delivered end-to-end database and data integration solutions aligned with business and infrastructure requirements.",
        ],
        "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel",
    )
    add_role(
        doc,
        "Data Warehouse DBA",
        "BAC Credomatic",
        "November 2017 - January 2019",
        "San Jose, Costa Rica",
        [
            "Developed and maintained ETL pipelines loading multiple source systems into an enterprise data warehouse.",
            "Defined and optimized table, index, and view structures for reliable, high-performance analytical workloads.",
            "Supported large datasets and reporting needs by improving structures used for analysis and business decision-making.",
        ],
        "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure",
    )

    add_heading(doc, "ADDITIONAL SQL SERVER EXPERIENCE")
    for item in [
        "Bosal, DBA and BI Consultant: designed the first stage of the main data warehouse, generated senior-management reports, and shared database improvement practices with development teams.",
        "Xetux Solutions, Database Manager: defined database management policies, improved database performance and security, and created a data lake for centralized sales reporting.",
        "ACH Cloud Services, SQL and BI Consultant: administered and developed company databases, created billing dashboards, and documented the data environment.",
        "EducaTablet, Microsoft SQL Server DBA: administered databases, improved database process response times, and trained development teams in database programming practices.",
        "VIGEOSOFT, Microsoft SQL Server DBA: modeled, designed, configured, and programmed OLTP databases and maintained data dictionary documentation.",
        "Optica Caroni, Developer Analyst: delivered SQL Server, SSIS, SSRS, .NET, and VB6 solutions for production workflows and management reporting.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "EDUCATION AND CERTIFICATIONS")
    for item in [
        "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000 - 2004)",
        "English Certificate - Universidad Central de Venezuela / Microsoft (2013 - 2015)",
        "SQL Admin Part 1; Analyzing and Visualizing Data with Power BI; Complete Data Science Training with Python for Data Analysis; Dataiku Core Designer",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "LANGUAGES")
    doc.add_paragraph("Spanish: Native | English: Full Professional")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()

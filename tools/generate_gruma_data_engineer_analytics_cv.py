from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Work\CVs")
SOURCE = ROOT / "Base" / "Pedro_Gutierrez_CV.docx"
OUTPUT = ROOT / "Output" / "Pedro_Gutierrez_GRUMA_Data_Engineer_Analytics_CV.docx"

BLUE = "1F4E79"
DARK = "222222"
GRAY = "5B5B5B"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def write(paragraph, text, size=9.1, bold=False, italic=False, color=None, font="Aptos"):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "BDD7EE")
    borders.append(bottom)


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    write(p, text.upper(), 10.2, True, color=BLUE, font="Aptos Display")
    add_rule(p)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.4)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, text)


def skill_line(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    write(p, f"{label}: ", 9, True)
    write(p, text, 9)


def role(doc, title, company, dates, place, points, technologies, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    write(p, f"{title} | {company}", 10, True, color=DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    write(p, f"{dates} | {place}", 8.5, italic=True, color=GRAY)

    for point in points:
        bullet(doc, point)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write(p, "Tecnologías: ", 8.6, True, color=DARK)
    write(p, technologies, 8.6, color=GRAY)


doc = Document(SOURCE)
clear_body(doc)
page = doc.sections[0]
page.top_margin = Inches(0.48)
page.bottom_margin = Inches(0.48)
page.left_margin = Inches(0.62)
page.right_margin = Inches(0.62)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9.1)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0.5)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 17, True, color=BLUE, font="Aptos Display")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "Ingeniero de Datos y BI | SQL Server, ETL/ELT, Power BI y Snowflake", 10, True, color="404040")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
write(p, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.5, color="4C4C4C")

section_heading(doc, "Resumen Profesional")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.05
p.paragraph_format.space_after = Pt(4)
write(
    p,
    "Ingeniero de Datos, desarrollador Snowflake y especialista SQL Server con más de 15 años de experiencia diseñando, integrando, transformando y optimizando soluciones de datos empresariales. Experiencia en ETL/ELT, SSIS, SQL Server, T-SQL, Snowflake, dbt, Data Warehouse, modelado dimensional Kimball, capas silver/gold, modelos semánticos, Power BI, SSAS, SSRS, Python, PowerShell y Azure. He construido pipelines y datasets analíticos, automatizado controles de calidad, optimizado procesos de almacenamiento y apoyado iniciativas de analítica e inteligencia artificial con Snowflake Cortex, MetricFlow y herramientas de desarrollo asistido por IA. Capacidad para traducir necesidades de negocio en soluciones escalables, confiables y documentadas.",
)

section_heading(doc, "Competencias Técnicas")
skill_line(doc, "Integración y transformación", "ETL/ELT, SSIS, dbt, pipelines, integración de múltiples fuentes, automatización, validaciones, manejo de errores y soporte productivo")
skill_line(doc, "Bases de datos", "SQL Server, T-SQL, Snowflake SQL, PostgreSQL, MySQL, stored procedures, consultas complejas, índices y optimización de rendimiento")
skill_line(doc, "Arquitectura y modelado", "Data Warehouse, modelado dimensional Kimball, modelos tabulares y semánticos, capas silver/gold, MetricFlow y datasets reutilizables")
skill_line(doc, "BI y analítica", "Power BI, SSAS, SSRS, Excel, preparación de datos, indicadores, análisis de tendencias, reporting y visualización")
skill_line(doc, "Automatización e IA", "Python, Pandas, NumPy, Matplotlib, PowerShell, Snowflake Cortex, Cursor y revisión automatizada de pull requests")
skill_line(doc, "Calidad y operación", "controles de calidad, reconciliación, detección de anomalías, monitoreo, trazabilidad, documentación e investigación de incidentes")

section_heading(doc, "Experiencia Profesional")
role(
    doc,
    "Snowflake Developer / Data Engineer",
    "ServiceTitan",
    "Septiembre 2024 - Actualidad",
    "Remoto / Costa Rica",
    [
        "Migro lógica de reportes desarrollada en C# hacia Snowflake SQL para habilitar procesos analíticos escalables y mantenibles.",
        "Creo y mantengo modelos dbt en capas silver y gold, mejorando la organización, reutilización y confiabilidad de los datos consumidos por reporting.",
        "Optimicé cargas Snowflake y dbt, reduciendo el tiempo de procesamiento del data warehouse en 20% durante una fase inicial de mejora.",
        "Desarrollé una prueba de concepto de modelado Data Warehouse con metodología Kimball para definir estándares y mejorar consultas y procesamiento.",
        "Apoyo la capa semántica de MetricFlow y flujos analíticos habilitados por Snowflake Cortex para organizar métricas y mantener definiciones consistentes.",
    ],
    "Snowflake, Snowflake SQL, dbt, MetricFlow, Snowflake Cortex, Kimball, Data Warehouse, SQL, Git, Cursor",
)

role(
    doc,
    "Data Engineer",
    "Health Catalyst",
    "Mayo 2021 - Septiembre 2024",
    "San Jose, Costa Rica",
    [
        "Diseñé e implementé aplicaciones ETL personalizadas que redujeron el procesamiento manual hasta en 40%.",
        "Construí y mantuve flujos de datos utilizando SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI y Excel.",
        "Automaticé análisis exploratorio, validaciones y detección de anomalías con Pandas, NumPy y Matplotlib, mejorando la precisión de validación en 30%.",
        "Preparé, limpié, transformé y modelé datasets para reporting, dashboards y toma de decisiones operativas.",
    ],
    "SQL Server, SSIS, Power BI, Python, Pandas, NumPy, Matplotlib, PowerShell, Snowflake, Excel",
)

role(
    doc,
    "SQL Developer",
    "Intertec International",
    "Febrero 2019 - Abril 2021",
    "San Jose, Costa Rica",
    [
        "Desarrollé y optimicé stored procedures, consultas SQL y flujos de datos para lógica de negocio y necesidades analíticas.",
        "Consolidé múltiples fuentes en datasets listos para análisis mediante procesos ETL optimizados.",
        "Mejoré la precisión de datos en más de 25% y reduje tiempos de ejecución hasta en 50% mediante validaciones, manejo de errores, índices y tuning SQL.",
    ],
    "SQL Server, T-SQL, SSIS, Salesforce, MySQL",
)

role(
    doc,
    "DBA / SQL Developer",
    "EL Tiempo",
    "Septiembre 2020 - Noviembre 2020",
    "Colombia",
    [
        "Lideré la planificación y ejecución de actividades de migración para diez bases SQL Server, manteniendo integridad de datos y continuidad operativa.",
        "Coordiné riesgos, validaciones, despliegues y transición con equipos multifuncionales.",
    ],
    "SQL Server, SSIS, Azure, SSRS",
    page_break_before=True,
)

role(
    doc,
    "DBA / SQL and BI Consultant",
    "Gold Data Networks",
    "Enero 2016 - Febrero 2020",
    "Ciudad de Panamá, Panamá",
    [
        "Diseñé bases relacionales, esquemas, tablas, stored procedures y objetos para aplicaciones operativas y cargas de reporting.",
        "Implementé soluciones de bases de datos, integración y BI alineadas con necesidades de negocio e infraestructura.",
    ],
    "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel",
)

role(
    doc,
    "Data Warehouse DBA",
    "BAC Credomatic",
    "Noviembre 2017 - Enero 2019",
    "San Jose, Costa Rica",
    [
        "Desarrollé y mantuve pipelines ETL para integrar múltiples fuentes en el Data Warehouse empresarial.",
        "Definí y optimicé tablas, índices y vistas para cargas analíticas confiables y de alto rendimiento.",
        "Analicé grandes volúmenes de datos para identificar tendencias y patrones útiles para la toma de decisiones.",
    ],
    "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure",
)

section_heading(doc, "Experiencia Adicional")
bullet(doc, "Bosal, DBA and BI Consultant: diseñé la primera etapa del Data Warehouse principal y generé reportes y presentaciones para la gerencia.")
bullet(doc, "Xetux Solutions, Database Manager: definí políticas de administración, mejoré rendimiento y seguridad, y creé un data lake centralizado para ventas.")
bullet(doc, "ACH Cloud Services y EducaTablet: administré bases de datos, desarrollé soluciones de reporting, documenté ambientes y capacité equipos técnicos.")
bullet(doc, "VIGEOSOFT y Optica Caroni: diseñé bases OLTP, desarrollé procesos SQL y aplicaciones, y entregué reportes operativos.")

section_heading(doc, "Educación e Idiomas")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
write(p, "Analista de Sistemas, Informática - IUT Dr. Federico Rivero Palacio (2000 - 2004).")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
write(p, "Formación adicional: Analyzing and Visualizing Data with Power BI; Python for Data Analysis; Dataiku Core Designer; SQL Admin Part 1.")
p = doc.add_paragraph()
write(p, "Español: Nativo | Inglés: Competencia profesional completa")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "GRUMA Data Engineer and Analytics CV"
doc.core_properties.subject = "SQL Server, ETL ELT, Data Warehouse, Power BI, Snowflake, Python and AI-assisted analytics"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)

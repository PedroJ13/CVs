from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_GIGA_IT_Senior_SQL_Server_DBA_Data_Analytics_CV.docx"


def write(paragraph, value, size=9.2, bold=False, color="26333F"):
    run = paragraph.add_run(value)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def section(doc, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    write(p, label.upper(), 10.5, True, "174B70")


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.7)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, value)


def role(doc, title, company, dates, place, points, technologies, break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = break_before
    write(p, f"{title} | {company}", 10, True, "17365D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    write(p, f"{dates} | {place}", 8.7, color="596878")
    for point in points:
        bullet(doc, point)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write(p, "Tecnologías: ", 8.6, True, "596878")
    write(p, technologies, 8.6, color="596878")


doc = Document()
page = doc.sections[0]
page.top_margin = Inches(0.52)
page.bottom_margin = Inches(0.52)
page.left_margin = Inches(0.65)
page.right_margin = Inches(0.65)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 16, True, "17365D")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "Senior SQL Server DBA | Automatización y análisis de datos", 10.1, True, "174B70")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
write(p, "San José, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.6, color="596878")

section(doc, "Perfil profesional")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.04
write(p, "Especialista en SQL Server, desarrollo SQL e ingeniería de datos con más de 15 años de experiencia en administración, optimización y evolución de plataformas de datos empresariales. Dominio de T-SQL, procedimientos almacenados, consultas complejas, índices, rendimiento, migraciones, integridad de datos y procesos ETL. Experiencia automatizando flujos con SSIS, Python y PowerShell, fortaleciendo prácticas de seguridad de datos y preparando información para Power BI, reportes operativos y toma de decisiones. Acostumbrado a trabajar de forma remota con equipos técnicos y de negocio.")

section(doc, "Competencias técnicas")
for label, value in [
    ("SQL Server y T-SQL", "Administración y desarrollo, procedimientos almacenados, consultas complejas, vistas, índices y optimización de rendimiento."),
    ("Automatización", "SSIS, ETL/ELT, PowerShell, Python, flujos de validación, procesamiento y actualización de datos."),
    ("Seguridad y calidad", "Políticas de administración y seguridad de datos, integridad, validaciones, manejo de errores y documentación técnica."),
    ("Análisis y BI", "Power BI, SSAS, SSRS, Excel, modelado de datos, datasets analíticos, tendencias y visualización para decisiones."),
    ("Otras plataformas", "Snowflake SQL, dbt, PostgreSQL, MySQL, Azure e integraciones con Salesforce."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    write(p, f"{label}: ", 9.1, True)
    write(p, value, 9.1)

section(doc, "Experiencia profesional")
role(doc, "Snowflake Developer / Data Engineer", "ServiceTitan", "Sep 2024 - Actualidad", "Remoto / Costa Rica", [
    "Migra lógica de reportes desde C# hacia Snowflake SQL y mantiene modelos dbt reutilizables para análisis y reporting.",
    "Optimiza cargas SQL y dbt, reduciendo en 20% el tiempo de procesamiento del data warehouse en una primera fase.",
    "Desarrolla capas silver y gold y apoya la capa semántica de MetricFlow para mantener definiciones analíticas consistentes.",
], "Snowflake SQL, dbt, MetricFlow, Snowflake Cortex, Kimball, Git")
role(doc, "Data Engineer", "SMASH Costa Rica", "May 2021 - Sep 2024", "San José, Costa Rica", [
    "Diseñó y mantuvo aplicaciones ETL y flujos productivos con SQL Server, SSIS, Python y PowerShell, reduciendo trabajo manual hasta en 40%.",
    "Automatizó análisis y validaciones para detectar anomalías y mejorar la exactitud de datos en 30%.",
    "Preparó y modeló datasets para Power BI, reportes empresariales y decisiones operativas.",
], "SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI, Excel")
role(doc, "SQL Developer", "Intertec International", "Feb 2019 - Abr 2021", "San José, Costa Rica", [
    "Desarrolló y optimizó procedimientos almacenados, consultas T-SQL y flujos de base de datos para lógica empresarial y análisis.",
    "Redujo tiempos de ejecución hasta en 50% mediante optimización SQL, estrategias de índices y análisis de rendimiento.",
    "Mejoró la exactitud de datos en más de 25% mediante validaciones y manejo de errores en procesos críticos.",
], "SQL Server, T-SQL, SSIS, Salesforce, MySQL")
role(doc, "DBA / SQL Developer", "EL Tiempo", "Sep 2020 - Nov 2020", "Colombia", [
    "Lideró la migración de diez bases SQL Server, coordinando validación, riesgos, integridad de datos y continuidad operativa.",
], "SQL Server, SSIS, Azure, SSRS", break_before=True)
role(doc, "DBA / SQL & BI Consultant", "Gold Data Networks", "Ene 2016 - Feb 2020", "Ciudad de Panamá, Panamá", [
    "Diseñó bases relacionales, tablas, objetos SQL y procedimientos almacenados para aplicaciones y cargas de reportes.",
    "Implementó soluciones integrales de base de datos, integración y BI alineadas con necesidades empresariales.",
], "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel")
role(doc, "Data Warehouse DBA", "BAC Credomatic", "Nov 2017 - Ene 2019", "San José, Costa Rica", [
    "Desarrolló y mantuvo pipelines ETL desde múltiples fuentes hacia el data warehouse empresarial.",
    "Diseñó y optimizó tablas, índices y vistas para cargas analíticas y análisis de grandes conjuntos de datos.",
], "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure")

section(doc, "Experiencia adicional")
bullet(doc, "Xetux Solutions, Database Manager: definió políticas de administración, mejoró el rendimiento y la seguridad de datos, y creó un data lake de ventas.")
bullet(doc, "EducaTablet, SQL Server DBA: administró bases empresariales, mejoró tiempos de respuesta y capacitó a desarrolladores en buenas prácticas SQL.")
bullet(doc, "VIGEOSOFT, SQL Server DBA: modeló y configuró bases OLTP y documentó cambios estructurales mediante diccionarios de datos.")
bullet(doc, "Bosal y ACH Cloud Services: diseñó componentes de data warehouse, creó reportes ejecutivos y documentó entornos de datos.")

section(doc, "Formación e idiomas")
p = doc.add_paragraph()
write(p, "Analista de Sistemas, Informática - IUT Dr. Federico Rivero Palacio (2000-2004).")
p = doc.add_paragraph()
write(p, "Formación adicional: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI; Python for Data Analysis.")
p = doc.add_paragraph()
write(p, "Español: nativo | Inglés: dominio profesional.")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "GIGA IT Senior SQL Server DBA Data Analytics CV"
doc.save(OUTPUT)
print(OUTPUT)

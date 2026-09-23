from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_TribuESK_Lead_Database_Data_Platform_Engineer_CV.docx"


def write(paragraph, value, size=9.2, bold=False, color="26333F"):
    run = paragraph.add_run(value)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def section(doc, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    write(p, label.upper(), 10.5, True, "174B70")


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.7)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, value)


def role(doc, title, company, dates, place, points, technologies, break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = break_before
    write(p, f"{title} | {company}", 10, True, "17365D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    write(p, f"{dates} | {place}", 8.7, color="596878")
    for point in points:
        bullet(doc, point)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write(p, "Technologies: ", 8.6, True, "596878")
    write(p, technologies, 8.6, color="596878")


doc = Document()
page = doc.sections[0]
page.top_margin = Inches(0.52)
page.bottom_margin = Inches(0.52)
page.left_margin = Inches(0.65)
page.right_margin = Inches(0.65)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 16, True, "17365D")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "Lead Database & Data Platform Engineer | SQL Server, Snowflake and MySQL", 10.1, True, "174B70")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
write(p, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.6, color="596878")

section(doc, "Professional Summary")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.04
write(p, "Database and data platform professional with 15+ years of experience designing, administering, optimizing, and modernizing enterprise data solutions. Strong background in Microsoft SQL Server, advanced T-SQL, MySQL, PostgreSQL, Snowflake, database migrations, performance tuning, data integrity, security practices, and data warehouse engineering. Experienced leading database migration activities, establishing database management policies, improving production data workflows, and guiding development teams on database programming practices. Works independently across remote and cross-functional teams and uses Snowflake Cortex, Cursor, and AI-assisted pull-request review to improve development, documentation, and quality.")

section(doc, "Core Expertise")
for label, value in [
    ("Database platforms", "Microsoft SQL Server, Snowflake, MySQL, PostgreSQL; relational design, schemas, tables, views, indexes, and stored procedures."),
    ("Performance and reliability", "Advanced T-SQL, query tuning, indexing strategies, workload optimization, troubleshooting, validation, and data integrity."),
    ("Platform engineering", "Database migrations, ETL/ELT, SSIS, dbt, dimensional modeling, data warehouse design, and production data support."),
    ("Standards and leadership", "Database management policies, security improvements, technical documentation, stakeholder coordination, developer training, and independent ownership."),
    ("Automation and AI", "Python, PowerShell, Git-based workflows, Cursor, Snowflake Cortex, and AI-assisted PR review."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    write(p, f"{label}: ", 9.1, True)
    write(p, value, 9.1)

section(doc, "Professional Experience")
role(doc, "Snowflake Developer / Data Engineer", "ServiceTitan", "Sep 2024 - Present", "Remote / Costa Rica", [
    "Own Snowflake SQL and dbt transformations that convert reporting logic into scalable, reusable analytical models.",
    "Optimize Snowflake and dbt workloads, reducing data warehouse processing time by 20% during an initial improvement phase.",
    "Created a Kimball-based warehouse modeling proof of concept to improve data structure, reuse, and query performance.",
    "Support silver and gold data layers, MetricFlow semantic models, and Snowflake Cortex use cases within Git-based development workflows.",
], "Snowflake, Snowflake SQL, dbt, MetricFlow, Kimball, Snowflake Cortex, Cursor, Git")
role(doc, "Data Engineer", "SMASH Costa Rica", "May 2021 - Sep 2024", "San Jose, Costa Rica", [
    "Designed and maintained production data workflows across SQL Server, SSIS, Snowflake, Python, PowerShell, Power BI, and Excel.",
    "Built client-specific ETL applications that reduced manual processing time by up to 40%.",
    "Automated anomaly detection and validation workflows, improving data validation accuracy by 30%.",
], "SQL Server, SSIS, Snowflake, Python, PowerShell, Power BI, Pandas")
role(doc, "SQL Developer", "Intertec International", "Feb 2019 - Apr 2021", "San Jose, Costa Rica", [
    "Developed and optimized advanced T-SQL queries, stored procedures, and database workflows supporting production business logic.",
    "Reduced query execution times by up to 50% through performance analysis, SQL tuning, and indexing strategies.",
    "Integrated SQL Server, Salesforce, and MySQL data while improving data accuracy by more than 25% through validation and error handling.",
], "SQL Server, T-SQL, SSIS, Salesforce, MySQL")
role(doc, "DBA / SQL Developer", "EL Tiempo", "Sep 2020 - Nov 2020", "Colombia", [
    "Led the planning and execution of migration activities for ten SQL Server databases while maintaining data integrity and operational continuity.",
    "Coordinated risk identification, validation, deployment timing, and transition activities with cross-functional teams.",
], "SQL Server, SSIS, Azure, SSRS", break_before=True)
role(doc, "DBA / SQL and BI Consultant", "Gold Data Networks", "Jan 2016 - Feb 2020", "Panama City, Panama", [
    "Designed relational databases, schemas, tables, stored procedures, and supporting objects for operational applications and reporting.",
    "Delivered end-to-end database, integration, and BI solutions aligned with business and infrastructure requirements.",
], "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel")
role(doc, "Data Warehouse DBA", "BAC Credomatic", "Nov 2017 - Jan 2019", "San Jose, Costa Rica", [
    "Developed and maintained ETL pipelines that loaded multiple sources into an enterprise data warehouse.",
    "Designed and optimized tables, indexes, and views for reliable, high-performance analytical workloads.",
], "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure")

section(doc, "Additional Leadership and Database Experience")
bullet(doc, "Xetux Solutions, Database Manager: defined database management policies, improved performance and data security, and created a centralized sales data lake.")
bullet(doc, "EducaTablet, SQL Server DBA: administered production databases, improved response times, and trained development teams in database programming practices.")
bullet(doc, "VIGEOSOFT, SQL Server DBA: modeled, designed, configured, and programmed OLTP databases and documented structural changes.")
bullet(doc, "Bosal and ACH Cloud Services: designed data warehouse components, administered databases, created management reports, and documented data environments.")

section(doc, "Education and Languages")
p = doc.add_paragraph()
write(p, "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000-2004).")
p = doc.add_paragraph()
write(p, "Additional training: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI; Python for Data Analysis.")
p = doc.add_paragraph()
write(p, "Spanish: Native | English: Full professional proficiency.")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "TribuESK Lead Database and Data Platform Engineer CV"
doc.save(OUTPUT)
print(OUTPUT)

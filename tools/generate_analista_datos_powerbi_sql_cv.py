from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import generate_fullstack_principal_sql_dba_cv as layout


SOURCE = layout.ROOT / "Base" / "Pedro_Gutierrez_CV.docx"
OUTPUT = layout.ROOT / "Output" / "Pedro_Gutierrez_Analista_Datos_PowerBI_SQL_CV.docx"


def add_skills_table(doc):
    rows = [
        (
            "Análisis y visualización",
            "Power BI, análisis exploratorio y descriptivo, dashboards, reportes, indicadores, identificación de tendencias, presentación de resultados y soporte a usuarios",
        ),
        (
            "SQL y preparación",
            "SQL avanzado, T-SQL, consultas complejas, vistas, procedimientos almacenados, integración, limpieza, transformación, validación y consolidación de datos",
        ),
        (
            "Modelado de datos",
            "Modelado relacional y dimensional, metodología Kimball, hechos y dimensiones, Data Warehouse, Data Lake, modelos semánticos y conjuntos de datos reutilizables",
        ),
        (
            "Herramientas analíticas",
            "Excel, Power BI, SSRS, SSAS, Python, Pandas, NumPy, Matplotlib, Snowflake, dbt, MetricFlow y SQL Server",
        ),
        (
            "Calidad y colaboración",
            "Controles de calidad, conciliación, detección de anomalías, documentación, levantamiento de requerimientos y comunicación con áreas técnicas y de negocio",
        ),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    widths = [Inches(1.72), Inches(5.68)]
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for index, text in enumerate(("Área", "Tecnologías y conocimientos")):
        cell = table.rows[0].cells[index]
        cell.text = ""
        layout.set_cell_shading(cell, layout.BLUE)
        layout.set_cell_borders(cell)
        layout.set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        layout.write(p, text, 8.7, True, color="FFFFFF")

    for row_index, (area, details) in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate((area, details)):
            cell = cells[index]
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            layout.set_cell_borders(cell)
            layout.set_cell_margins(cell)
            if row_index % 2 == 0:
                layout.set_cell_shading(cell, "F3F7FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            layout.write(p, value, 8.35, bold=(index == 0))


def experiencia(doc, cargo, empresa, fechas, ubicacion, puntos, tecnologias, salto=False):
    layout.role(doc, cargo, empresa, fechas, ubicacion, puntos, tecnologias, page_break_before=salto)


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)
    layout.clear_body(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.43)
    section.bottom_margin = Inches(0.43)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.1)
    normal.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0.5)
    layout.write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 18, True, color=layout.BLUE, font="Aptos Display")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    layout.write(p, "Analista de Datos | Power BI, SQL y Modelado de Datos", 9.8, True, color="404040")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    layout.write(p, "San José, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.5, color="4C4C4C")

    layout.section_heading(doc, "Perfil profesional")
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.04
    p.paragraph_format.space_after = Pt(4)
    layout.write(
        p,
        "Profesional de datos con amplia experiencia en análisis, preparación, integración y modelado de información para reportes, dashboards y toma de decisiones. Dominio de SQL y experiencia práctica con Power BI, Excel, SQL Server, Snowflake, SSIS, SSRS, SSAS, Python y herramientas de calidad de datos. Capacidad para recopilar información de múltiples fuentes, crear conjuntos de datos confiables, identificar tendencias y anomalías, optimizar consultas y comunicar resultados a equipos técnicos y de negocio. Experiencia desarrollando Data Warehouses, Data Lakes, modelos dimensionales y capas semánticas para facilitar el acceso consistente a indicadores e información operativa.",
    )

    layout.section_heading(doc, "Competencias técnicas")
    add_skills_table(doc)

    layout.section_heading(doc, "Experiencia profesional")
    experiencia(
        doc,
        "Snowflake Developer / Data Engineer",
        "ServiceTitan",
        "Septiembre 2024 - Actualidad",
        "Remoto / Costa Rica",
        [
            "Migro lógica de reportes desarrollada en C# hacia Snowflake SQL, creando transformaciones mantenibles para análisis y consumo de información.",
            "Desarrollo y mantengo modelos dbt en capas silver y gold, mejorando la organización, reutilización y confiabilidad de los conjuntos de datos analíticos.",
            "Optimicé cargas de Snowflake SQL y dbt, reduciendo en 20% el tiempo de procesamiento del Data Warehouse durante una fase inicial de mejora.",
            "Desarrollé una prueba de concepto de modelado dimensional basada en Kimball y apoyo modelos semánticos en MetricFlow para mantener definiciones consistentes de métricas.",
        ],
        "Snowflake, Snowflake SQL, dbt, MetricFlow, Kimball, Data Warehouse, SQL y Git",
    )

    experiencia(
        doc,
        "Data Engineer",
        "Health Catalyst",
        "Mayo 2021 - Septiembre 2024",
        "San José, Costa Rica",
        [
            "Preparé, limpié, transformé y modelé datos para dashboards de Power BI, reportes de negocio y análisis operativo.",
            "Diseñé aplicaciones ETL personalizadas que redujeron hasta en 40% el procesamiento manual de información.",
            "Automaticé análisis exploratorios y controles de validación con Python y Pandas, mejorando en 30% la precisión de la validación de datos.",
            "Integré y mantuve flujos de datos con SQL Server, SSIS, Snowflake, PowerShell, Power BI y Excel.",
        ],
        "Power BI, Excel, SQL Server, T-SQL, SSIS, Snowflake, Python, Pandas, NumPy, Matplotlib y PowerShell",
    )

    experiencia(
        doc,
        "SQL Developer",
        "Intertec International",
        "Febrero 2019 - Abril 2021",
        "San José, Costa Rica",
        [
            "Desarrollé y optimicé consultas SQL, procedimientos almacenados y flujos de datos para análisis y lógica de negocio.",
            "Consolidé fuentes de SQL Server, Salesforce y MySQL en conjuntos de datos preparados para análisis mediante procesos ETL.",
            "Mejoré la precisión de los datos en más de 25% mediante validaciones y manejo de errores, y reduje tiempos de consulta hasta en 50% mediante optimización SQL.",
        ],
        "SQL Server, T-SQL, SSIS, Salesforce y MySQL",
        salto=True,
    )

    experiencia(
        doc,
        "DBA / SQL and BI Consultant",
        "Gold Data Networks",
        "Enero 2016 - Febrero 2020",
        "Ciudad de Panamá, Panamá",
        [
            "Diseñé bases de datos relacionales, estructuras de reportes y procesos de integración para requerimientos operativos y de inteligencia de negocios.",
            "Entregué soluciones con SQL Server, SSIS, SSRS, Power BI, PostgreSQL y Excel alineadas con las necesidades de usuarios y áreas de negocio.",
        ],
        "Power BI, Excel, SQL Server, T-SQL, SSIS, SSRS y PostgreSQL",
    )

    experiencia(
        doc,
        "Data Warehouse DBA",
        "BAC Credomatic",
        "Noviembre 2017 - Enero 2019",
        "San José, Costa Rica",
        [
            "Desarrollé y mantuve procesos ETL que integraban múltiples sistemas fuente en el Data Warehouse empresarial.",
            "Diseñé y optimicé tablas, índices, vistas y estructuras analíticas para reportes confiables y de alto rendimiento.",
            "Analicé grandes conjuntos de datos para identificar tendencias y patrones que apoyaban la toma de decisiones.",
        ],
        "Power BI, Excel, SQL Server, SSIS, SSAS, SSRS y Azure",
    )

    layout.section_heading(doc, "Experiencia adicional relevante")
    layout.bullet(doc, "Bosal, DBA and BI Consultant: diseñé la primera etapa del Data Warehouse principal y elaboré reportes y presentaciones para la alta gerencia.", 8.7)
    layout.bullet(doc, "Xetux Solutions, Database Manager: establecí políticas de gestión de datos y construí un Data Lake centralizado de ventas para análisis y reportes.", 8.7)
    layout.bullet(doc, "ACH Cloud Services, SQL and BI Consultant: desarrollé reportes de facturación y documenté el entorno de datos de la empresa.", 8.7)
    layout.bullet(doc, "Optica Caroni, Developer Analyst: implementé proyectos operativos y desarrollé reportes y dashboards para la administración.", 8.7)

    layout.section_heading(doc, "Educación, formación e idiomas")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    layout.write(p, "Analista de Sistemas, Informática - IUT Dr. Federico Rivero Palacio (2000 - 2004)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    layout.write(p, "Formación: Analyzing and Visualizing Data with Power BI; Python for Data Analysis; Dataiku Core Designer; SQL Admin Part 1")
    p = doc.add_paragraph()
    layout.write(p, "Español: Nativo | Inglés: Profesional")

    doc.core_properties.author = "Pedro Javier Gutierrez Armas"
    doc.core_properties.title = "Currículum Analista de Datos Power BI SQL"
    doc.core_properties.subject = "Currículum adaptado para análisis de datos, Power BI, SQL y modelado de información"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

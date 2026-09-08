from docx import Document

import generate_senior_microsoft_sql_server_dba_partner_cv as source


DOCX_PATH = source.OUTPUT_DIR / "Pedro_Gutierrez_Prod_Support_Data_Engineer_TSQL_CV.docx"


def replace_paragraph(paragraph, text, bold=False):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    run.bold = bold


def build_document():
    source.DOCX_PATH = DOCX_PATH
    source.build_document()

    doc = Document(DOCX_PATH)
    for paragraph in doc.paragraphs:
        if paragraph.text == "Senior Microsoft SQL Server DBA | T-SQL Performance | Data Platforms":
            replace_paragraph(
                paragraph,
                "Production Support Data Engineer | T-SQL | SQL Server | SSIS",
                bold=True,
            )
        elif paragraph.text.startswith("Senior Microsoft SQL Server DBA, SQL Developer, and Data Engineer"):
            replace_paragraph(
                paragraph,
                "Production Support Data Engineer and SQL Server specialist with 15+ years of experience developing, maintaining, optimizing, and supporting SQL-based data integration, ETL, reporting, and database solutions. Strong hands-on experience with Microsoft SQL Server, T-SQL, SSIS, stored procedures, indexing, slow-query tuning, data validation, PowerShell, Power BI, and production troubleshooting. Proven ability to analyze live data issues, improve query execution time, maintain reliable data workflows, support database migrations and releases, document technical processes, and collaborate with technical teams and business stakeholders in English and Spanish.",
            )
        elif paragraph.text.startswith("Troubleshot data and database workflow issues while consolidating"):
            replace_paragraph(
                paragraph,
                "Troubleshot production data and database workflow issues while consolidating multiple source systems into analytics-ready datasets.",
            )

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            area = row.cells[0].text
            if area == "SQL Server Administration":
                row.cells[1].text = "Microsoft SQL Server, production support, database maintenance, migrations, operational continuity, issue resolution"
            elif area == "Data Architecture and BI":
                row.cells[1].text = "SSIS, ETL/ELT, SQL Server data integration, data warehousing, data modeling, SSAS, SSRS, Power BI"
            elif area == "Data Operations":
                row.cells[1].text = "Production workflow support, data integrity, validation checks, error handling, process documentation, stakeholder coordination"

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()

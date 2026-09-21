from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_Analista_Estructura_SQL_Server_CV.docx"


def run(p, value, size=9.2, bold=False, color="25313D"):
    r = p.add_run(value)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run(p, title.upper(), 10.5, True, "174B70")


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.8)
    p.paragraph_format.line_spacing = 1.02
    run(p, value)


def role(doc, title, company, dates, location, points, tech, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    run(p, f"{title} | {company}", 10, True, "17365D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run(p, f"{dates} | {location}", 8.7, color="596878")
    for point in points:
        bullet(doc, point)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, "Tecnologías: ", 8.6, True, "596878")
    run(p, tech, 8.6, color="596878")


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.52)
sec.bottom_margin = Inches(0.52)
sec.left_margin = Inches(0.65)
sec.right_margin = Inches(0.65)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(9.2)
doc.styles["Normal"].paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
run(p, "PEDRO JAVIER GUTIERREZ ARMAS", 16, True, "17365D")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
run(p, "Analista de bases de datos | SQL Server, T-SQL y documentación técnica", 10.1, True, "174B70")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
run(p, "San José, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.6, color="596878")

heading(doc, "Perfil profesional")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.04
run(p, "Especialista en SQL Server y desarrollo de datos con más de 15 años de experiencia en diseño, administración, análisis y optimización de bases de datos empresariales. Experiencia en estructuras relacionales, tablas, vistas, índices, procedimientos almacenados, lógica de negocio en SQL y documentación de cambios mediante diccionarios de datos. Ha trabajado en migración de bases de datos, integración ETL y análisis de consultas para entender dependencias, conservar la integridad de los datos y apoyar la evolución de plataformas existentes.")

heading(doc, "Competencias técnicas")
for label, value in [
    ("SQL Server y T-SQL", "Diseño y análisis de tablas, vistas, índices, objetos de base de datos, consultas complejas y procedimientos almacenados."),
    ("Modelado y documentación", "Modelado relacional y OLTP, estructuras de datos, diccionarios de datos, documentación de procesos y cambios."),
    ("Diagnóstico y evolución", "Análisis de lógica SQL, rendimiento de consultas, migraciones, validación, integridad y continuidad operativa."),
    ("Integración y BI", "SSIS, SSAS, SSRS, Power BI, ETL, data warehouse, Snowflake, dbt, PostgreSQL y MySQL."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, label + ": ", 9.1, True)
    run(p, value, 9.1)

heading(doc, "Experiencia profesional")
role(doc, "Snowflake Developer / Data Engineer", "ServiceTitan", "Sep 2024 - Actualidad", "Remoto / Costa Rica", [
    "Migra lógica de reportes implementada en C# a Snowflake SQL y mantiene modelos dbt en capas silver y gold para consumo analítico.",
    "Creó una prueba de concepto de modelado de data warehouse con metodología Kimball para estandarizar estructuras y mejorar consultas y procesamiento.",
    "Optimiza cargas SQL y dbt; en una primera fase redujo en 20% el tiempo de procesamiento del data warehouse.",
], "Snowflake SQL, dbt, modelado de datos, Kimball, MetricFlow" )
role(doc, "Data Engineer", "SMASH Costa Rica", "May 2021 - Sep 2024", "San José, Costa Rica", [
    "Diseñó aplicaciones ETL según requisitos de clientes y mantuvo flujos de datos con SQL Server y SSIS, reduciendo trabajo manual hasta en 40%.",
    "Preparó, transformó y modeló conjuntos de datos para reportes y tableros de Power BI; automatizó controles de validación para detectar anomalías.",
], "SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI")
role(doc, "SQL Developer", "Intertec International", "Feb 2019 - Abr 2021", "San José, Costa Rica", [
    "Desarrolló y optimizó procedimientos almacenados, consultas SQL y flujos de base de datos que soportaban lógica de negocio y necesidades analíticas.",
    "Integró fuentes de datos diferentes mediante ETL, validaciones y manejo de errores; mejoró la exactitud de datos en más de 25%.",
    "Redujo tiempos de ejecución de consultas hasta en 50% mediante optimización SQL, estrategias de índices y análisis de rendimiento.",
], "SQL Server, T-SQL, SSIS, Salesforce, MySQL", page_break_before=True)
role(doc, "DBA / SQL Developer", "EL Tiempo", "Sep 2020 - Nov 2020", "Colombia", [
    "Lideró la planificación y ejecución de la migración de diez bases de datos, con validación de datos, identificación de riesgos y continuidad operativa.",
], "SQL Server, SSIS, Azure, SSRS")
role(doc, "DBA / SQL & BI Consultant", "Gold Data Networks", "Ene 2016 - Feb 2020", "Ciudad de Panamá, Panamá", [
    "Construyó bases de datos relacionales y estructuras de tablas desde cero para aplicaciones web y cargas de reportes.",
    "Desarrolló objetos de base de datos y procedimientos almacenados para apoyar lógica de aplicaciones, integración y rendimiento.",
], "SQL Server, SSIS, SSRS, PostgreSQL, Power BI")
role(doc, "Data Warehouse DBA", "BAC Credomatic", "Nov 2017 - Ene 2019", "San José, Costa Rica", [
    "Desarrolló y mantuvo procesos ETL de múltiples fuentes hacia el data warehouse; definió y optimizó tablas, índices y vistas para cargas analíticas.",
], "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure")

heading(doc, "Experiencia adicional relevante")
bullet(doc, "VIGEOSOFT (2014-2015): modeló, diseñó y configuró bases OLTP; documentó cambios estructurales y de procesos mediante diccionarios de datos.")
bullet(doc, "ACH Cloud Services (2015-2016): administró y desarrolló bases de datos y documentó el entorno de datos de la empresa.")
bullet(doc, "Bosal, Xetux Solutions, EducaTablet y Optica Caroni: administración SQL Server, diseño de soluciones de datos, reportes y mejora de procesos.")

heading(doc, "Formación e idiomas")
p = doc.add_paragraph()
run(p, "Analista de Sistemas, Informática - IUT Dr. Federico Rivero Palacio (2000-2004).", 9.1)
p = doc.add_paragraph()
run(p, "Formación adicional: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI.", 9.1)
p = doc.add_paragraph()
run(p, "Español: nativo | Inglés: dominio profesional.", 9.1)

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "Analista de estructura de bases de datos SQL Server"
doc.save(OUTPUT)
print(OUTPUT)

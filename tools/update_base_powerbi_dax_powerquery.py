from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(r"C:\Work\CVs")
GENERAL = ROOT / "Base" / "Pedro_Gutierrez_CV.docx"
DBA = ROOT / "Base" / "Pedro_Gutierrez_CV_SQL_Server_DBA.docx"

HEALTH_POWER_BI_BULLET = (
    "Worked for six months directly within a Power BI team, developing and supporting "
    "dashboards, data models, DAX measures, Power Query transformations, and reporting workflows."
)


def replace_text(paragraph, old, new):
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    return False


def insert_after(reference, text):
    new_p = OxmlElement("w:p")
    if reference._p.pPr is not None:
        new_p.append(deepcopy(reference._p.pPr))
    reference._p.addnext(new_p)
    paragraph = Paragraph(new_p, reference._parent)
    paragraph.add_run(text)
    return paragraph


def add_health_power_bi_bullet(doc):
    if any(p.text == HEALTH_POWER_BI_BULLET for p in doc.paragraphs):
        return

    title_index = next(
        i
        for i, p in enumerate(doc.paragraphs)
        if "Health Catalyst" in p.text
    )
    technology_index = next(
        i
        for i in range(title_index + 1, len(doc.paragraphs))
        if doc.paragraphs[i].text.startswith("Technologies:")
    )
    reference = doc.paragraphs[technology_index - 1]
    insert_after(reference, HEALTH_POWER_BI_BULLET)


def update_general(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if text.startswith("Data Engineer, Snowflake Developer"):
            replace_text(paragraph, "Power BI, Python", "Power BI, DAX, Power Query, Python")
        elif text.startswith("Microsoft BI Stack:"):
            replace_text(paragraph, "Power BI, Excel", "Power BI, DAX, Power Query, Excel")
        elif text.startswith("Built and maintained data workflows"):
            replace_text(
                paragraph,
                "Snowflake, Power BI, and Excel",
                "Snowflake, Power BI, DAX, Power Query, and Excel",
            )
        elif text.startswith("Prepared, cleaned, transformed, and modeled datasets"):
            if "DAX measures" in text:
                replace_text(
                    paragraph,
                    "Prepared, cleaned, transformed, and modeled datasets for business reporting, Power BI dashboards, DAX measures, and Power Query transformations, and operational decision-making.",
                    "Prepared, cleaned, transformed, and modeled datasets and reporting components for Power BI dashboards and operational decision-making, including DAX measures and Power Query transformations.",
                )
            else:
                replace_text(
                    paragraph,
                    "Prepared, cleaned, transformed, and modeled datasets for business reporting, Power BI dashboards, and operational decision-making.",
                    "Prepared, cleaned, transformed, and modeled datasets and reporting components for Power BI dashboards and operational decision-making, including DAX measures and Power Query transformations.",
                )
        elif text.startswith("Technologies:") and "Power BI" in text and "DAX" not in text:
            replace_text(paragraph, "Power BI", "Power BI, DAX, Power Query")

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "Power BI" in paragraph.text and "DAX" not in paragraph.text:
                        replace_text(paragraph, "Power BI", "Power BI, DAX, Power Query")

    add_health_power_bi_bullet(doc)


def update_dba(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Technologies:") and "Power BI" in paragraph.text and "DAX" not in paragraph.text:
            replace_text(paragraph, "Power BI", "Power BI, DAX, Power Query")
    add_health_power_bi_bullet(doc)


def main():
    general = Document(GENERAL)
    update_general(general)
    general.save(GENERAL)

    dba = Document(DBA)
    update_dba(dba)
    dba.save(DBA)

    print(GENERAL)
    print(DBA)


if __name__ == "__main__":
    main()

from docx import Document
from docx.shared import Pt

import generate_fullstack_principal_sql_dba_cv as source


DOCX_PATH = source.OUTPUT_DIR / "Pedro_Gutierrez_Agentic_Commerce_Analytics_Engineer_CV.docx"


def replace_paragraph(paragraph, text, bold=False):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    run.bold = bold


def build_document():
    source.DOCX_PATH = DOCX_PATH
    source.build_document()

    doc = Document(DOCX_PATH)
    for paragraph in doc.paragraphs:
        if paragraph.text == "Principal SQL Database Administrator | SQL Server DBA | T-SQL Performance Tuning":
            replace_paragraph(
                paragraph,
                "Analytics Engineer | Snowflake | dbt | Semantic Layer | AI-Enabled Data",
                bold=True,
            )
        elif paragraph.text.startswith("Principal SQL Server DBA, SQL Developer, and Data Engineer"):
            replace_paragraph(
                paragraph,
                "Analytics Engineer, Snowflake Developer, and Data Engineer with 15+ years of experience transforming complex source data into reliable models for analytics, reporting, and business decision-making. Current hands-on work includes Snowflake SQL, dbt models across silver and gold layers, Kimball modeling, MetricFlow semantic-layer development, Snowflake Cortex, performance optimization, and migration of reporting logic from C# into maintainable SQL. Strong background in SQL, Python, ETL/ELT, data quality, validation, documentation, stakeholder collaboration, and AI-assisted development using Cursor and automated pull request review workflows.",
            )
        elif paragraph.text.startswith("Support MetricFlow semantic layer development and Snowflake Cortex"):
            replace_paragraph(
                paragraph,
                "Support MetricFlow semantic-layer development and Snowflake Cortex use cases, improving metric consistency and maintaining clear analytical definitions for downstream users and AI-enabled workflows.",
            )
        elif paragraph.text == "ADDITIONAL SQL SERVER DBA AND DATABASE EXPERIENCE":
            replace_paragraph(paragraph, "ADDITIONAL DATA AND DATABASE EXPERIENCE", bold=True)

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            area = row.cells[0].text
            if area == "SQL Server Administration":
                row.cells[0].text = "Snowflake and dbt"
                row.cells[1].text = "Snowflake SQL, dbt models, silver and gold layers, warehouse optimization, reusable analytical assets"
            elif area == "Performance Engineering":
                row.cells[0].text = "Analytics Engineering"
                row.cells[1].text = "Advanced SQL, query optimization, transformation logic, data warehousing, dimensional modeling, Kimball practices"
            elif area == "Database Development":
                row.cells[0].text = "Semantic Layer"
                row.cells[1].text = "MetricFlow, metric organization, consistent analytical definitions, Snowflake Cortex, AI-enabled analytics"
            elif area == "Data Engineering and BI":
                row.cells[0].text = "Pipelines and Quality"
                row.cells[1].text = "ETL/ELT, SSIS, Python, data transformation, validation checks, anomaly detection, Power BI, reporting datasets"
            elif area == "Automation and Delivery":
                row.cells[0].text = "Development Workflow"
                row.cells[1].text = "Python, PowerShell, Git exposure, documentation, Cursor, AI-assisted pull request review, cross-functional delivery"

            if area != "Area":
                for run in row.cells[0].paragraphs[0].runs:
                    run.bold = True
                for cell in row.cells:
                    for cell_paragraph in cell.paragraphs:
                        for run in cell_paragraph.runs:
                            run.font.size = Pt(8.5)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()

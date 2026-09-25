from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Work\CVs")
SOURCE = ROOT / "Base" / "Pedro_Gutierrez_CV_SQL_Server_DBA.docx"
OUTPUT = ROOT / "Output" / "Pedro_Gutierrez_DWConsulware_DBA_SQL_Server_Senior_CV.docx"

BLUE = "1F4E79"
DARK = "222222"
GRAY = "5B5B5B"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def write(paragraph, text, size=9.1, bold=False, italic=False, color=None, font="Aptos"):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "BDD7EE")
    borders.append(bottom)


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    write(p, text.upper(), 10.2, True, color=BLUE, font="Aptos Display")
    add_rule(p)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.4)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, text)


def skill_line(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    write(p, f"{label}: ", 9, True)
    write(p, text, 9)


def role(doc, title, company, dates, place, points, technologies, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    write(p, f"{title} | {company}", 10, True, color=DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    write(p, f"{dates} | {place}", 8.5, italic=True, color=GRAY)

    for point in points:
        bullet(doc, point)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write(p, "Tecnologías: ", 8.6, True, color=DARK)
    write(p, technologies, 8.6, color=GRAY)


doc = Document(SOURCE)
clear_body(doc)
page = doc.sections[0]
page.top_margin = Inches(0.48)
page.bottom_margin = Inches(0.48)
page.left_margin = Inches(0.62)
page.right_margin = Inches(0.62)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9.1)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0.5)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 17, True, color=BLUE, font="Aptos Display")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "DBA SQL Server Senior | Migraciones, Rendimiento y Soporte Productivo", 10, True, color="404040")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
write(p, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.5, color="4C4C4C")

section_heading(doc, "Resumen Profesional")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.05
p.paragraph_format.space_after = Pt(4)
write(
    p,
    "DBA SQL Server, ingeniero de bases de datos y desarrollador T-SQL con más de 15 años de experiencia administrando, desarrollando, optimizando y soportando plataformas de datos empresariales. Experiencia sólida en SQL Server, T-SQL avanzado, migraciones de bases de datos, respaldo y recuperación, seguridad, monitoreo, SQL Agent, resolución de incidentes, ajuste de rendimiento, SSIS, PowerShell y continuidad operativa. Capacidad demostrada para liderar actividades de migración, ejecutar validaciones, coordinar despliegues controlados, documentar procedimientos y colaborar con equipos de desarrollo, infraestructura, QA, soporte y negocio.",
)

section_heading(doc, "Competencias Técnicas")
skill_line(doc, "Administración SQL Server", "operación de ambientes productivos, acceso y seguridad, backup/restore, mantenimiento, SQL Agent, monitoreo, resolución de incidentes y soporte de disponibilidad")
skill_line(doc, "Migraciones", "planificación, identificación de riesgos, validación de datos, coordinación de despliegues, continuidad operativa, recuperación y verificación posterior")
skill_line(doc, "T-SQL y rendimiento", "stored procedures, funciones, consultas complejas, índices, análisis de ejecución, procesos de larga duración, bloqueos y optimización set-based")
skill_line(doc, "Automatización e integración", "PowerShell, SSIS, ETL/ELT, Python, Git, scripts de despliegue y rollback, herramientas de diagnóstico y documentación operativa")
skill_line(doc, "Plataformas", "SQL Server, AWS RDS, Azure, MySQL, PostgreSQL, Snowflake, SSAS, SSRS y Power BI")

section_heading(doc, "Experiencia Profesional")
role(
    doc,
    "SQL Server DBA / Database Engineer",
    "MWR Life",
    "Septiembre 2024 - Actualidad",
    "Remoto / Costa Rica",
    [
        "Administro y doy soporte a bases SQL Server alojadas en AWS RDS en ambientes de Producción, Desarrollo y Staging, incluyendo accesos, conectividad, backup/restore, mantenimiento y atención de incidentes.",
        "Diseñé un modelo de acceso basado en roles, mínimo privilegio y acceso privilegiado JIT; audité logins, usuarios, permisos y procedimientos posteriores a restauraciones.",
        "Desarrollo stored procedures, funciones, scripts de despliegue y rollback, procedimientos de diagnóstico y utilidades de soporte productivo.",
        "Lideré una recuperación controlada de datos e integración InEvent, validando respaldos, recuperando cientos de registros, corrigiendo procedimientos y verificando el despliegue.",
        "Analizo consultas de larga duración, índices, patrones de ejecución y procesos basados en cursores para implementar alternativas set-based y mejorar el rendimiento.",
    ],
    "SQL Server, AWS RDS, T-SQL, SSMS, DBeaver, SQL Agent, Database Mail, PowerShell, JSON/OpenJSON, Git, RBAC, JIT, backup/restore, monitoreo",
)

role(
    doc,
    "DBA / SQL Developer",
    "Health Catalyst",
    "Mayo 2021 - Septiembre 2024",
    "San Jose, Costa Rica",
    [
        "Desarrollé y soporté flujos de datos productivos utilizando SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI y Excel.",
        "Diseñé aplicaciones ETL personalizadas que redujeron el procesamiento manual hasta en 40%.",
        "Automaticé validaciones y detección de anomalías, mejorando la precisión de validación de datos en 30%.",
    ],
    "SQL Server, T-SQL, SSIS, Snowflake, Python, PowerShell, Power BI, Pandas, Excel",
)

role(
    doc,
    "SQL Developer",
    "Intertec International",
    "Febrero 2019 - Abril 2021",
    "San Jose, Costa Rica",
    [
        "Desarrollé y optimicé consultas T-SQL, stored procedures y flujos de bases de datos para lógica de negocio y analítica.",
        "Reduje los tiempos de ejecución hasta en 50% mediante optimización SQL, estrategias de índices y análisis de rendimiento.",
        "Integré datos de SQL Server, Salesforce y MySQL, mejorando la precisión en más de 25% mediante validaciones y manejo de errores.",
    ],
    "SQL Server, T-SQL, SSIS, Salesforce, MySQL",
)

role(
    doc,
    "DBA / SQL Developer",
    "EL Tiempo",
    "Septiembre 2020 - Noviembre 2020",
    "Colombia",
    [
        "Lideré la planificación y ejecución de actividades de migración para diez bases SQL Server, completando las fases dentro del alcance y plazo establecidos.",
        "Coordiné riesgos, validación de datos, controles de integridad, tiempos de despliegue y continuidad operativa con equipos multifuncionales.",
    ],
    "SQL Server, SSIS, Azure, SSRS",
    page_break_before=True,
)

role(
    doc,
    "DBA / SQL and BI Consultant",
    "Gold Data Networks",
    "Enero 2016 - Febrero 2020",
    "Ciudad de Panamá, Panamá",
    [
        "Diseñé bases relacionales, esquemas, tablas, stored procedures y objetos para aplicaciones operativas y cargas de reporting.",
        "Entregué soluciones de bases de datos, integración y BI alineadas con requisitos de rendimiento e infraestructura.",
    ],
    "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel",
)

role(
    doc,
    "Data Warehouse DBA",
    "BAC Credomatic",
    "Noviembre 2017 - Enero 2019",
    "San Jose, Costa Rica",
    [
        "Desarrollé y mantuve procesos ETL para cargar múltiples fuentes al data warehouse empresarial.",
        "Diseñé y optimicé tablas, índices y vistas para cargas analíticas confiables y de alto rendimiento.",
    ],
    "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure",
)

section_heading(doc, "Experiencia Adicional")
bullet(doc, "Bosal, DBA and BI Consultant: diseñé la primera etapa del data warehouse principal y apoyé soluciones SQL Server, SSIS, Azure, MySQL y Power BI.")
bullet(doc, "Xetux Solutions, Database Manager: definí políticas de administración, mejoré rendimiento y seguridad, y creé un data lake centralizado para ventas.")
bullet(doc, "ACH Cloud Services y EducaTablet: administré bases empresariales, desarrollé procesos SQL y reporting, documenté ambientes y capacité equipos de desarrollo.")
bullet(doc, "VIGEOSOFT y Optica Caroni: diseñé bases OLTP, desarrollé procesos SQL y aplicaciones, y entregué soluciones de reporting operativo.")

section_heading(doc, "Educación e Idiomas")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
write(p, "Analista de Sistemas, Informática - IUT Dr. Federico Rivero Palacio (2000 - 2004).")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
write(p, "Formación adicional: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI; Python for Data Analysis.")
p = doc.add_paragraph()
write(p, "Español: Nativo | Inglés: Competencia profesional completa")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "DBA SQL Server Senior CV"
doc.core_properties.subject = "SQL Server, migraciones, rendimiento, backup y recuperación, PowerShell y soporte productivo"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)

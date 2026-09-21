from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_Metrica_Data_Engineer_Markets_CV.docx"


def write(paragraph, value, size=9.2, bold=False, color="26333F"):
    run = paragraph.add_run(value)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def section(doc, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    write(p, value.upper(), 10.5, True, "174B70")


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.7)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, value)


def role(doc, title, company, dates, place, points, technologies, break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = break_before
    write(p, f"{title} | {company}", 10, True, "17365D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    write(p, f"{dates} | {place}", 8.7, color="596878")
    for point in points:
        bullet(doc, point)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write(p, "Tecnologías: ", 8.6, True, "596878")
    write(p, technologies, 8.6, color="596878")


doc = Document()
page = doc.sections[0]
page.top_margin = Inches(0.52)
page.bottom_margin = Inches(0.52)
page.left_margin = Inches(0.65)
page.right_margin = Inches(0.65)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 16, True, "17365D")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "Data Engineer | SQL, ETL, modelado y calidad de datos", 10.1, True, "174B70")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
write(p, "San José, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.6, color="596878")

section(doc, "Perfil profesional")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.04
write(p, "Ingeniero de datos y especialista SQL con más de 15 años de experiencia en desarrollo, integración, automatización, modelado y validación de datos. Experiencia con SQL Server, T-SQL, SSIS, Snowflake, dbt, Python y PowerShell; diseño de procesos ETL/ELT y estructuras de data warehouse. Ha preparado información de múltiples fuentes para análisis y reportes, optimizado procesos productivos y colaborado con equipos técnicos y de negocio. Su experiencia en banca incluye el desarrollo de pipelines y estructuras analíticas para el data warehouse de BAC Credomatic.")

section(doc, "Competencias técnicas")
for label, value in [
    ("Ingeniería y modelado", "ETL/ELT, ingesta, transformación, data warehouse, Kimball, modelos silver y gold, tablas, vistas e índices."),
    ("SQL y automatización", "SQL Server, T-SQL, Snowflake SQL, SSIS, dbt, Python, PowerShell, consultas complejas y optimización."),
    ("Validación y operación", "Pruebas y controles de calidad de datos, validación de migraciones, detección de anomalías, manejo de errores y documentación."),
    ("Analítica", "Power BI, SSAS, SSRS, Excel, modelado de datasets para reportes y seguimiento de resultados."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    write(p, f"{label}: ", 9.1, True)
    write(p, value, 9.1)

section(doc, "Experiencia profesional")
role(doc, "Snowflake Developer / Data Engineer", "ServiceTitan", "Sep 2024 - Actualidad", "Remoto / Costa Rica", [
    "Migra lógica de reportes desde C# hacia Snowflake SQL y mantiene modelos dbt en capas silver y gold para datos analíticos reutilizables.",
    "Creó una prueba de concepto de data warehouse con metodología Kimball para definir estándares de modelado y mejorar el procesamiento.",
    "Optimiza procesos SQL y dbt; redujo en 20% el tiempo de procesamiento en una primera fase de optimización.",
], "Snowflake SQL, dbt, Kimball, MetricFlow, Snowflake Cortex")
role(doc, "Data Engineer", "SMASH Costa Rica", "May 2021 - Sep 2024", "San José, Costa Rica", [
    "Diseñó aplicaciones ETL personalizadas y flujos de datos con SQL Server, SSIS, Python y PowerShell; redujo trabajo manual hasta en 40%.",
    "Desarrolló análisis exploratorio y validaciones automatizadas para detectar anomalías y mejorar la exactitud de validación en 30%.",
    "Preparó, depuró, transformó y modeló datasets para reportes de negocio y tableros Power BI.",
], "SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI, Pandas")
role(doc, "SQL Developer", "Intertec International", "Feb 2019 - Abr 2021", "San José, Costa Rica", [
    "Desarrolló y optimizó procedimientos almacenados, consultas SQL y flujos de base de datos para lógica de negocio y análisis.",
    "Integró fuentes de datos mediante ETL y controles de validación y errores; mejoró la exactitud de datos en más de 25%.",
    "Redujo tiempos de ejecución de consultas hasta en 50% mediante optimización SQL e índices.",
], "SQL Server, SSIS, Salesforce, MySQL")
role(doc, "DBA / SQL Developer", "EL Tiempo", "Sep 2020 - Nov 2020", "Colombia", [
    "Lideró la migración de diez bases de datos, con identificación de riesgos, validación de datos y continuidad de operación.",
], "SQL Server, SSIS, Azure, SSRS", break_before=True)
role(doc, "DBA / SQL & BI Consultant", "Gold Data Networks", "Ene 2016 - Feb 2020", "Ciudad de Panamá, Panamá", [
    "Construyó bases relacionales, tablas, objetos SQL y procesos de datos integrados con aplicaciones y necesidades de reportes.",
], "SQL Server, SSIS, SSRS, PostgreSQL, Power BI")
role(doc, "Data Warehouse DBA", "BAC Credomatic", "Nov 2017 - Ene 2019", "San José, Costa Rica", [
    "Desarrolló y mantuvo pipelines ETL desde múltiples fuentes hacia el data warehouse de la entidad bancaria.",
    "Definió y optimizó tablas, índices y vistas para cargas analíticas y análisis de grandes conjuntos de datos.",
], "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure")

section(doc, "Experiencia adicional")
bullet(doc, "Bosal: diseñó la primera etapa del data warehouse principal y generó reportes para la dirección.")
bullet(doc, "Xetux Solutions: creó un data lake para consolidar información de ventas y apoyar análisis centralizado.")
bullet(doc, "ACH Cloud Services y VIGEOSOFT: administración de bases de datos, reportes, modelado y documentación técnica.")

section(doc, "Formación e idiomas")
p = doc.add_paragraph()
write(p, "Analista de Sistemas, Informática - IUT Dr. Federico Rivero Palacio (2000-2004).")
p = doc.add_paragraph()
write(p, "Formación adicional: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI; Python for Data Analysis.")
p = doc.add_paragraph()
write(p, "Español: nativo | Inglés: dominio profesional.")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "Metrica Data Engineer Markets CV"
doc.save(OUTPUT)
print(OUTPUT)

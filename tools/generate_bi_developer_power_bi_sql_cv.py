from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_BI_Developer_Power_BI_SQL_CV.docx"


def add_text(paragraph, value, size=9.2, bold=False, color="25313D"):
    run = paragraph.add_run(value)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def section(doc, title):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    add_text(paragraph, title.upper(), 10.5, True, "174B70")


def bullet(doc, value):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(1.7)
    paragraph.paragraph_format.line_spacing = 1.02
    paragraph.paragraph_format.keep_together = True
    add_text(paragraph, value)


def role(doc, title, company, dates, location, bullets, technologies, page_break_before=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    add_text(paragraph, f"{title} | {company}", 10, True, "17365D")

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    add_text(paragraph, f"{dates} | {location}", 8.7, color="596878")

    for point in bullets:
        bullet(doc, point)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    add_text(paragraph, "Technologies: ", 8.6, True, "596878")
    add_text(paragraph, technologies, 8.6, color="596878")


doc = Document()
page = doc.sections[0]
page.top_margin = Inches(0.52)
page.bottom_margin = Inches(0.52)
page.left_margin = Inches(0.65)
page.right_margin = Inches(0.65)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(2)

paragraph = doc.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_after = Pt(1)
add_text(paragraph, "PEDRO JAVIER GUTIERREZ ARMAS", 16, True, "17365D")

paragraph = doc.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_after = Pt(1)
add_text(paragraph, "BI Developer | Power BI, SQL Server and Data Modeling", 10.1, True, "174B70")

paragraph = doc.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_after = Pt(6)
add_text(paragraph, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.6, color="596878")

section(doc, "Professional Summary")
paragraph = doc.add_paragraph()
paragraph.paragraph_format.line_spacing = 1.04
add_text(paragraph, "Data and BI professional with 15+ years of experience in SQL Server development, reporting, ETL, data warehousing, and analytics. Hands-on experience preparing and modeling data for Power BI dashboards, building SQL reporting datasets, integrating disparate sources, and improving data quality and query performance. Background in dimensional modeling, SSAS and SSRS, with current work on dbt models and the MetricFlow semantic layer. Experienced working with business and technical stakeholders to translate reporting needs into reliable data assets and documented processes.")

section(doc, "Relevant Skills")
skills = [
    ("BI and reporting", "Power BI dashboards and reporting datasets, SSRS, SSAS, Excel, management reporting."),
    ("SQL and transformation", "SQL Server, T-SQL, stored procedures, complex queries, SSIS, ETL/ELT, Snowflake SQL."),
    ("Data modeling", "Kimball practices, dimensional data warehouse design, relational structures, dbt silver and gold models, MetricFlow semantic layer."),
    ("Quality and performance", "Data validation, anomaly detection, query tuning, indexing, warehouse processing optimization."),
    ("Collaboration", "Requirements translation, stakeholder communication, technical documentation, data dictionaries, Git-based workflows."),
]
for label, value in skills:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1.6)
    add_text(paragraph, f"{label}: ", 9.1, True)
    add_text(paragraph, value, 9.1)

section(doc, "Professional Experience")
role(doc, "Snowflake Developer / Data Engineer", "ServiceTitan", "Sep 2024 - Present", "Remote / Costa Rica", [
    "Migrate C# reporting logic into Snowflake SQL, creating maintainable data transformations for downstream reporting and analytics.",
    "Build and maintain dbt models in silver and gold layers, improving reuse and consistency of analytical datasets.",
    "Created a Kimball-based data warehouse modeling proof of concept and support MetricFlow semantic-layer development for consistent metric definitions.",
    "Optimized Snowflake SQL and dbt processing, reducing data warehouse processing time by 20% in an initial optimization phase.",
], "Snowflake, Snowflake SQL, dbt, Kimball, MetricFlow, Snowflake Cortex, Git-based workflows")

role(doc, "Data Engineer", "SMASH Costa Rica", "May 2021 - Sep 2024", "San Jose, Costa Rica", [
    "Designed client-specific ETL applications and data workflows using SQL Server and SSIS, reducing manual processing by up to 40%.",
    "Prepared, cleaned, transformed, and modeled datasets for Power BI dashboards and operational reporting.",
    "Built Python-based exploratory analysis and validation tools to detect anomalies and improve data validation accuracy by 30%.",
], "SQL Server, SSIS, Power BI, Python, PowerShell, Snowflake, Pandas, Excel")

role(doc, "SQL Developer", "Intertec International", "Feb 2019 - Apr 2021", "San Jose, Costa Rica", [
    "Developed and optimized SQL queries, stored procedures, and database workflows for complex business logic and analytics.",
    "Consolidated multiple sources into analytics-ready datasets with ETL, validation checks, and error handling, improving data accuracy by more than 25%.",
    "Reduced query execution times by up to 50% through SQL tuning and indexing strategies.",
], "SQL Server, T-SQL, SSIS, Salesforce, MySQL")

role(doc, "DBA / SQL Developer", "EL Tiempo", "Sep 2020 - Nov 2020", "Colombia", [
    "Led migration activities for 10 SQL Server databases, coordinating validation and data-integrity checks with cross-functional teams.",
], "SQL Server, SSIS, Azure, SSRS", page_break_before=True)

role(doc, "DBA / SQL and BI Consultant", "Gold Data Networks", "Jan 2016 - Feb 2020", "Panama City, Panama", [
    "Built relational database structures and reporting solutions integrated with existing client infrastructure.",
    "Developed database objects and stored procedures supporting application and BI workloads; worked with Power BI, SSRS, and Excel.",
], "SQL Server, T-SQL, SSIS, SSRS, PostgreSQL, Power BI, Excel")

role(doc, "Data Warehouse DBA", "BAC Credomatic", "Nov 2017 - Jan 2019", "San Jose, Costa Rica", [
    "Developed and maintained SSIS pipelines to bring data from multiple sources into the enterprise data warehouse.",
    "Defined and optimized tables, indexes, and views for high-performance analytical workloads and supported BI reporting.",
], "SQL Server, SSIS, SSAS, Power BI, SSRS, Azure")

section(doc, "Additional Relevant Experience")
bullet(doc, "Bosal: designed the first stage of its main data warehouse and delivered reports and presentations for senior management.")
bullet(doc, "ACH Cloud Services: created billing dashboards and reports and documented the company data environment.")
bullet(doc, "VIGEOSOFT: modeled OLTP databases and documented structural changes through data dictionaries and stakeholder coordination.")
bullet(doc, "Earlier roles also include SQL Server administration, data development, reporting, and team training at Xetux Solutions, EducaTablet, and Optica Caroni.")

section(doc, "Education and Training")
paragraph = doc.add_paragraph()
add_text(paragraph, "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000-2004).")
paragraph = doc.add_paragraph()
add_text(paragraph, "Analyzing and Visualizing Data with Power BI; SQL Admin Part 1; Dataiku Core Designer.")

section(doc, "Languages")
paragraph = doc.add_paragraph()
add_text(paragraph, "Spanish: Native | English: Full professional proficiency")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "BI Developer Power BI SQL CV"
doc.save(OUTPUT)
print(OUTPUT)

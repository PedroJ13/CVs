from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Work\CVs")
SOURCE = ROOT / "Base" / "Pedro_Gutierrez_CV.docx"
OUTPUT = ROOT / "Base" / "Pedro_Gutierrez_CV_SQL_Server_DBA.docx"

BLUE = "1F4E79"
DARK = "222222"
GRAY = "5B5B5B"
LIGHT_LINE = "BDD7EE"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="B7C9D6", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), LIGHT_LINE)
    borders.append(bottom)


def write(paragraph, text, size=9, bold=False, italic=False, color=None, font="Aptos"):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    write(p, text.upper(), 10, True, color=BLUE, font="Aptos Display")
    set_bottom_border(p)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, text)
    return p


def role(doc, title, company, dates, place, points, technologies, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    write(p, f"{title} | {company}", 10, True, color=DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    date_line = dates if not place else f"{dates} | {place}"
    write(p, date_line, 8.5, italic=True, color=GRAY)

    for point in points:
        bullet(doc, point)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    write(p, "Technologies: ", 8.6, True, color=DARK)
    write(p, technologies, 8.6, color=GRAY)


def add_strengths_table(doc):
    table = doc.add_table(rows=3, cols=3)
    table.autofit = False
    widths = [Inches(2.46)] * 3
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)

    headers = ["SQL Server and Security", "Performance and Reliability", "Data Platforms and Automation"]
    details = [
        "AWS RDS, T-SQL, RBAC, least privilege, JIT access, login and permission auditing",
        "Query tuning, indexes, execution analysis, monitoring, backup and restore, production support",
        "SSIS, Snowflake, dbt, MySQL, PostgreSQL, Python, PowerShell, Git and AI-assisted workflows",
    ]
    evidence = [
        "Standardized access controls and post-restore permission procedures across production and lower environments",
        "Reduced query execution times by up to 50% and led controlled recovery of hundreds of production records",
        "Reduced manual processing by up to 40% and improved data validation accuracy by 30%",
    ]
    for i, value in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        write(p, value, 8.6, True, color="FFFFFF")
    for row_index, values in ((1, details), (2, evidence)):
        for i, value in enumerate(values):
            cell = table.cell(row_index, i)
            set_cell_shading(cell, "F3F7FA" if row_index == 1 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            write(p, value, 8.2)
    return table


doc = Document(SOURCE)
clear_body(doc)

section = doc.sections[0]
section.top_margin = Inches(0.45)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0.5)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 18, True, color=BLUE, font="Aptos Display")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "Senior SQL Server DBA | Database Engineer | AWS RDS, Security, Performance and Production Support", 9.5, True, color="404040")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
write(p, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.5, color="4C4C4C")

section_heading(doc, "Professional Summary")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.05
p.paragraph_format.space_after = Pt(5)
write(
    p,
    "Senior SQL Server DBA, database engineer, and SQL developer with 15+ years of experience administering, securing, developing, optimizing, and supporting enterprise databases and data platforms. Hands-on experience with SQL Server on AWS RDS, advanced T-SQL, role-based access control, least-privilege and JIT access, backup and restore operations, production troubleshooting, performance tuning, SQL Agent monitoring, database migrations, SSIS, MySQL, PostgreSQL, Snowflake, PowerShell, and data warehousing. Proven ability to lead controlled recovery and deployment work, resolve complex security and connectivity issues, establish operational standards, document repeatable procedures, and collaborate with developers, QA, support, and business stakeholders.",
)

section_heading(doc, "Core Strengths")
add_strengths_table(doc)

section_heading(doc, "Technical Skills")
for item in [
    "SQL Server administration: AWS RDS, SSMS, DBeaver, database access, backup and restore, SQL Agent, Database Mail, monitoring, troubleshooting, and production support",
    "Database security: RBAC, least privilege, JIT privileged access, logins, users, roles, schema permissions, SID mapping, credential rotation, and access auditing",
    "SQL development and performance: advanced T-SQL, stored procedures, functions, deployment and rollback scripts, JSON/OpenJSON, indexing, execution analysis, long-running queries, and set-based optimization",
    "Data platforms and integration: SSIS, ETL/ELT, Snowflake, dbt, MySQL, PostgreSQL, Azure-based environments, data warehousing, and dimensional modeling",
    "Automation and engineering tools: PowerShell, Python, Git, Cursor, Snowflake Cortex, AI-assisted pull-request review, monitoring dashboards, and diagnostic utilities",
    "Operations and delivery: incident response, controlled deployments, data validation, recovery verification, runbooks, stakeholder coordination, and developer guidance",
]:
    bullet(doc, item)

section_heading(doc, "Professional Experience")
role(
    doc,
    "SQL Server DBA / Database Engineer",
    "MWR Life",
    "September 2024 - Present",
    "Remote / Costa Rica",
    [
        "Administer and support SQL Server databases hosted on AWS RDS across Production, Development, and Staging environments, including access, connectivity, backup and restore, maintenance, and production incident support.",
        "Designed a role-based access model using least-privilege principles, standardized usr_ and prv_ account prefixes, JIT privileged access, permission audits, and controlled post-restore access procedures for lower environments.",
        "Create and maintain T-SQL stored procedures, functions, deployment and rollback scripts, diagnostic procedures, and utilities while resolving authentication, SID mapping, SSL, firewall, and application connection issues.",
        "Led a controlled InEvent production data recovery and integration remediation effort, validating source backups, recovering hundreds of records, correcting stored procedures, and completing post-deployment verification.",
        "Analyze complex workloads, long-running queries, indexes, execution patterns, and cursor-based processes; develop set-based and configurable alternatives and validate equivalent results.",
        "Develop SQL monitoring dashboards and operational documentation for SQL Agent jobs, alerts, execution history, integrations, data-quality exceptions, restore procedures, JIT access, and incident response.",
    ],
    "SQL Server, AWS RDS, T-SQL, SSMS, DBeaver, SQL Agent, Database Mail, PowerShell, JSON/OpenJSON, Git, RBAC, JIT access, database security, backup/restore, monitoring, performance tuning",
)

role(
    doc,
    "DBA / SQL Developer",
    "SMASH Costa Rica",
    "May 2021 - September 2024",
    "San Jose, Costa Rica",
    [
        "Developed and supported SQL Server, SSIS, Snowflake, Python, and PowerShell data workflows for production and client-facing requirements.",
        "Designed customized ETL applications that reduced manual processing time by up to 40%.",
        "Automated validation and anomaly-detection routines, improving data validation accuracy by 30% and supporting reliable reporting operations.",
    ],
    "SQL Server, T-SQL, SSIS, Snowflake, Python, PowerShell, Power BI, Pandas, Excel",
    page_break_before=True,
)

role(
    doc,
    "SQL Developer",
    "Intertec International",
    "February 2019 - April 2021",
    "San Jose, Costa Rica",
    [
        "Developed and optimized advanced T-SQL queries, stored procedures, and database workflows supporting production business logic and analytics.",
        "Reduced query execution times by up to 50% through SQL tuning, indexing strategies, and performance analysis.",
        "Integrated SQL Server, Salesforce, and MySQL data while improving data accuracy by more than 25% through validation and error handling.",
    ],
    "SQL Server, T-SQL, SSIS, Salesforce, MySQL",
)

role(
    doc,
    "DBA / SQL Developer",
    "EL Tiempo",
    "September 2020 - November 2020",
    "Colombia",
    [
        "Led planning and execution of migration activities for ten SQL Server databases, delivering all phases on time and within scope.",
        "Coordinated migration risk analysis, data validation, integrity controls, deployment timing, and continuity with cross-functional teams.",
    ],
    "SQL Server, SSIS, Azure, SSRS",
)

role(
    doc,
    "DBA / SQL and BI Consultant",
    "Gold Data Networks",
    "January 2016 - February 2020",
    "Panama City, Panama",
    [
        "Designed relational databases, schemas, tables, stored procedures, and supporting objects for operational applications and reporting workloads.",
        "Delivered database, integration, and BI solutions aligned with application performance and infrastructure requirements.",
    ],
    "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel",
)

role(
    doc,
    "Data Warehouse DBA",
    "BAC Credomatic",
    "November 2017 - January 2019",
    "San Jose, Costa Rica",
    [
        "Developed and maintained ETL pipelines that loaded multiple source systems into the enterprise data warehouse.",
        "Designed and optimized tables, indexes, and views for reliable, high-performance analytical workloads.",
    ],
    "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure",
)

doc.add_section(WD_SECTION.NEW_PAGE)
section = doc.sections[-1]
section.top_margin = Inches(0.45)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)

section_heading(doc, "Additional Experience")
for title, company, dates, place, points, technologies in [
    ("DBA and BI Consultant", "Bosal", "March 2017 - March 2018", "Lummen, Belgium", ["Designed the first stage of the company's main data warehouse and generated reports for senior management.", "Shared database improvement practices with development teams and supported BI and reporting delivery."], "SQL Server, SSIS, Azure, MySQL, Power BI"),
    ("Database Manager", "Xetux Solutions", "August 2016 - October 2017", "Caracas, Venezuela", ["Defined database management policies and improved performance and data security across business areas.", "Created a centralized sales data lake for reporting and analysis."], "SQL Server, SSIS, SSRS, Azure"),
    ("SQL and BI Consultant", "ACH Cloud Services", "August 2015 - August 2016", "Caracas, Venezuela", ["Administered and developed databases supporting operational and reporting workflows.", "Created management dashboards and documented the company's data environment."], "SQL Server, SSIS, SSRS, Azure"),
    ("Microsoft SQL Server DBA", "EducaTablet", "February 2015 - August 2015", "Caracas, Venezuela", ["Administered and developed production databases and improved database process response times.", "Trained development teams on database programming practices."], "SQL Server, SSIS, SSRS, Azure"),
    ("Microsoft SQL Server DBA", "VIGEOSOFT", "August 2014 - January 2015", "Caracas, Venezuela", ["Modeled, designed, configured, and programmed OLTP databases for business applications.", "Documented process and structural changes through data dictionary practices."], "SQL Server, SSIS, SSRS"),
    ("Developer Analyst", "Optica Caroni C.A.", "September 2005 - May 2014", "Caracas, Venezuela", ["Implemented six projects for lens laboratory operations and improved production-related workflows.", "Adjusted application and database processes and created management reports and dashboards."], "SQL Server, SSIS, SSRS, .NET, VB6"),
]:
    role(doc, title, company, dates, place, points, technologies)

section_heading(doc, "Education and Certifications")
for item in [
    "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000 - 2004)",
    "English Certificate - Universidad Central de Venezuela / Microsoft (2013 - 2015)",
    "Analyzing and Visualizing Data with Power BI",
    "Complete Data Science Training with Python for Data Analysis",
    "R Programming A-Z: R for Data Science",
    "Dataiku Core Designer",
    "SQL Admin Part 1",
]:
    bullet(doc, item)

section_heading(doc, "Languages")
p = doc.add_paragraph()
write(p, "Spanish: Native | English: Full Professional")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "SQL Server DBA and Database Engineer Base CV"
doc.core_properties.subject = "Base resume for SQL Server DBA, SQL Developer, and database engineering roles"
doc.save(OUTPUT)
print(OUTPUT)

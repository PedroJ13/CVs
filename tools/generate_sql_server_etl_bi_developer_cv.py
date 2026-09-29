from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import generate_fullstack_principal_sql_dba_cv as layout


OUTPUT = layout.ROOT / "Output" / "Pedro_Gutierrez_SQL_Server_ETL_BI_Developer_CV.docx"


def add_skills_table(doc):
    rows = [
        (
            "SQL Server and T-SQL",
            "Production database administration and development, advanced T-SQL, stored procedures, functions, views, indexing, execution analysis, query optimization and troubleshooting",
        ),
        (
            "ETL and reporting",
            "SSIS, SSRS, ETL/ELT pipelines, multi-source integration, transformations, scheduling, validation, error handling, reporting datasets and data-quality controls",
        ),
        (
            "Automation",
            "PowerShell, Python, Pandas, deployment and rollback scripts, diagnostic procedures, monitoring utilities and repeatable operational runbooks",
        ),
        (
            "Data platforms",
            "SQL Server, AWS RDS, Snowflake, MySQL, PostgreSQL, Azure-based environments, SSAS, Power BI, Excel and data warehousing",
        ),
        (
            "Data architecture",
            "Relational modeling, normalization, OLTP design, dimensional modeling, schemas, tables, facts, dimensions, data integrity, security and access controls",
        ),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    widths = [Inches(1.72), Inches(5.68)]
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for index, text in enumerate(("Area", "Tools and practices")):
        cell = table.rows[0].cells[index]
        cell.text = ""
        layout.set_cell_shading(cell, layout.BLUE)
        layout.set_cell_borders(cell)
        layout.set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        layout.write(p, text, 8.7, True, color="FFFFFF")

    for row_index, (area, details) in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate((area, details)):
            cell = cells[index]
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            layout.set_cell_borders(cell)
            layout.set_cell_margins(cell)
            if row_index % 2 == 0:
                layout.set_cell_shading(cell, "F3F7FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            layout.write(p, value, 8.35, bold=(index == 0))


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(layout.SOURCE)
    layout.clear_body(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.43)
    section.bottom_margin = Inches(0.43)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.1)
    normal.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0.5)
    layout.write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 18, True, color=layout.BLUE, font="Aptos Display")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    layout.write(p, "SQL Server Database Developer and Administrator | ETL, T-SQL and BI", 9.8, True, color="404040")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    layout.write(p, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.5, color="4C4C4C")

    layout.section_heading(doc, "Professional Summary")
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.04
    p.paragraph_format.space_after = Pt(4)
    layout.write(
        p,
        "SQL Server database developer, administrator, and data engineer with 15+ years of experience supporting production databases, building ETL and reporting solutions, and optimizing data workloads. Strong hands-on background in advanced T-SQL, stored procedures, query tuning, indexing, SSIS, SSRS, PowerShell, Python, relational and dimensional modeling, data quality, and database operations. Experienced administering SQL Server on AWS RDS across production and non-production environments, integrating SQL Server with Snowflake, MySQL, PostgreSQL, Salesforce, Power BI, and Azure-based services, and translating business requirements into reliable data solutions.",
    )

    layout.section_heading(doc, "Technical Skills")
    add_skills_table(doc)

    layout.section_heading(doc, "Professional Experience")
    layout.role(
        doc,
        "SQL Server DBA / Database Engineer",
        "MWR Life",
        "September 2024 - Present",
        "Remote / Costa Rica",
        [
            "Administer and support SQL Server databases hosted on AWS RDS across Production, Development, and Staging, including access, connectivity, backup and restore, maintenance, monitoring, and incident response.",
            "Develop and maintain T-SQL stored procedures, functions, deployment and rollback scripts, diagnostic procedures, and production support utilities.",
            "Analyze long-running queries, indexes, execution patterns, query variants, and cursor-based processes; implement set-based alternatives and validate equivalent results.",
            "Built diagnostic tooling for event and webhook integrations to identify missing records, invalid dates, duplicate data, unmatched entities, and data-quality issues.",
            "Automate and document SQL Agent monitoring, execution history, alerts, restore procedures, data validation, controlled deployments, and operational support processes.",
        ],
        "SQL Server, AWS RDS, T-SQL, SSMS, DBeaver, SQL Agent, Database Mail, PowerShell, JSON/OpenJSON, Git, backup/restore, monitoring and performance tuning",
    )

    layout.role(
        doc,
        "DBA / SQL Developer",
        "Health Catalyst",
        "May 2021 - September 2024",
        "San Jose, Costa Rica",
        [
            "Developed and supported production data workflows using SQL Server, T-SQL, SSIS, Snowflake, Python, PowerShell, Power BI, and Excel.",
            "Designed customized ETL applications that reduced manual data processing time by up to 40%.",
            "Automated data validation and anomaly-detection routines with Python and Pandas, improving validation accuracy by 30% and supporting reliable reporting.",
            "Prepared, cleaned, transformed, and modeled data for dashboards, business reporting, and operational decision-making.",
        ],
        "SQL Server, T-SQL, SSIS, Snowflake, Python, PowerShell, Power BI, Pandas and Excel",
    )

    layout.role(
        doc,
        "SQL Developer",
        "Intertec International",
        "February 2019 - April 2021",
        "San Jose, Costa Rica",
        [
            "Developed and optimized advanced T-SQL queries, stored procedures, and database workflows supporting production business logic and analytics.",
            "Reduced query execution times by up to 50% through SQL tuning, indexing strategies, and performance analysis.",
            "Integrated SQL Server, Salesforce, and MySQL data through ETL processes while improving data accuracy by more than 25% through validation and error handling.",
        ],
        "SQL Server, T-SQL, SSIS, Salesforce and MySQL",
        page_break_before=True,
    )

    layout.role(
        doc,
        "DBA / SQL Developer",
        "EL Tiempo",
        "September 2020 - November 2020",
        "Colombia",
        [
            "Led planning and execution of migration activities for ten SQL Server databases, delivering all phases on time and within scope.",
            "Coordinated data validation, integrity controls, migration risk analysis, deployment timing, and operational continuity with cross-functional teams.",
        ],
        "SQL Server, SSIS, Azure and SSRS",
    )

    layout.role(
        doc,
        "DBA / SQL and BI Consultant",
        "Gold Data Networks",
        "January 2016 - February 2020",
        "Panama City, Panama",
        [
            "Designed relational databases, schemas, tables, stored procedures, and integration components for operational and reporting workloads.",
            "Delivered SQL Server, SSIS, SSRS, PostgreSQL, Power BI, and reporting solutions aligned with business and infrastructure requirements.",
        ],
        "SQL Server, SSIS, SSRS, PostgreSQL, Power BI and Excel",
    )

    layout.role(
        doc,
        "Data Warehouse DBA",
        "BAC Credomatic",
        "November 2017 - January 2019",
        "San Jose, Costa Rica",
        [
            "Developed and maintained SSIS ETL pipelines that loaded multiple source systems into the enterprise data warehouse.",
            "Designed and optimized tables, indexes, views, and analytical structures for reliable, high-performance reporting workloads.",
        ],
        "SQL Server, T-SQL, SSIS, SSAS, SSRS, Power BI and Azure",
    )

    layout.section_heading(doc, "Additional Database and BI Experience")
    layout.bullet(doc, "Bosal, DBA and BI Consultant: designed the first stage of the main data warehouse and delivered management reporting with SQL Server, SSIS, Azure, MySQL, and Power BI.", 8.7)
    layout.bullet(doc, "Xetux Solutions, Database Manager: defined database-management policies, improved performance and data security, and created a centralized sales data lake.", 8.7)
    layout.bullet(doc, "ACH Cloud Services, SQL and BI Consultant: administered databases, created billing-management reports, and documented the data environment using SQL Server, SSIS, and SSRS.", 8.7)
    layout.bullet(doc, "EducaTablet and VIGEOSOFT, Microsoft SQL Server DBA: administered production databases, improved response times, modeled OLTP systems, and trained development teams.", 8.7)
    layout.bullet(doc, "Optica Caroni, Developer Analyst: delivered operational database, integration, and reporting projects using SQL Server, SSIS, SSRS, .NET, and VB6.", 8.7)

    layout.section_heading(doc, "Education and Languages")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    layout.write(p, "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000 - 2004)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    layout.write(p, "Training: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI; Dataiku Core Designer; Python for Data Analysis")
    p = doc.add_paragraph()
    layout.write(p, "Spanish: Native | English: Full Professional")

    doc.core_properties.author = "Pedro Javier Gutierrez Armas"
    doc.core_properties.title = "SQL Server ETL and BI Developer Resume"
    doc.core_properties.subject = "Resume tailored for a SQL Server database development, administration, ETL, and BI role"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

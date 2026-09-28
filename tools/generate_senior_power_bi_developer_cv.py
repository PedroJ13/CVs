from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Work\CVs")
SOURCE = ROOT / "Base" / "Pedro_Gutierrez_CV.docx"
OUTPUT = ROOT / "Output" / "Pedro_Gutierrez_Senior_Power_BI_Developer_CV.docx"

BLUE = "1F4E79"
DARK = "222222"
GRAY = "5B5B5B"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def write(paragraph, text, size=9.1, bold=False, italic=False, color=None, font="Aptos"):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "BDD7EE")
    borders.append(bottom)


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    write(p, text.upper(), 10.2, True, color=BLUE, font="Aptos Display")
    add_rule(p)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.4)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, text)


def skill_line(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    write(p, f"{label}: ", 9, True)
    write(p, text, 9)


def role(doc, title, company, dates, place, points, technologies, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    write(p, f"{title} | {company}", 10, True, color=DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    write(p, f"{dates} | {place}", 8.5, italic=True, color=GRAY)

    for point in points:
        bullet(doc, point)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write(p, "Technologies: ", 8.6, True, color=DARK)
    write(p, technologies, 8.6, color=GRAY)


doc = Document(SOURCE)
clear_body(doc)
page = doc.sections[0]
page.top_margin = Inches(0.48)
page.bottom_margin = Inches(0.48)
page.left_margin = Inches(0.62)
page.right_margin = Inches(0.62)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9.1)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0.5)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 17, True, color=BLUE, font="Aptos Display")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "Senior Power BI Developer | SQL, Data Modeling and ETL", 10, True, color="404040")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
write(p, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.5, color="4C4C4C")

section_heading(doc, "Professional Summary")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.05
p.paragraph_format.space_after = Pt(4)
write(
    p,
    "Power BI, SQL, and data professional with 15+ years of experience delivering reporting, business intelligence, data integration, and data warehouse solutions. Power BI and enterprise reporting experience across roles since 2017, supported by deep expertise in SQL Server, advanced T-SQL, SSIS, SSAS, SSRS, Snowflake, dbt, dimensional modeling, semantic models, data quality, and performance optimization. Experienced gathering requirements, preparing and modeling large datasets, building reliable reporting layers, analyzing operational information, troubleshooting reporting issues, and communicating actionable insights to technical and business stakeholders.",
)

section_heading(doc, "Technical Skills")
skill_line(doc, "Business intelligence", "Power BI dashboards and reports, visual analysis, KPI reporting, SSRS, SSAS, Excel, report support, and performance improvement")
skill_line(doc, "SQL development", "advanced T-SQL, complex queries, views, stored procedures, functions, indexing, execution analysis, and query optimization")
skill_line(doc, "Data modeling", "dimensional modeling, Kimball methodology, Data Warehouse design, semantic models, MetricFlow, facts and dimensions, and reusable reporting datasets")
skill_line(doc, "ETL and integration", "SSIS, ETL/ELT pipelines, dbt, multi-source integration, transformations, scheduling, validation, and error handling")
skill_line(doc, "Data platforms", "SQL Server, Snowflake, PostgreSQL, MySQL, Azure-based environments, Salesforce integrations, and large analytical datasets")
skill_line(doc, "Quality and automation", "data reconciliation, anomaly detection, Python, Pandas, NumPy, Matplotlib, PowerShell, Git, documentation, and production support")

section_heading(doc, "Professional Experience")
role(
    doc,
    "Snowflake Developer / Data Engineer",
    "ServiceTitan",
    "September 2024 - Present",
    "Remote / Costa Rica",
    [
        "Migrate C# reporting logic into Snowflake SQL to provide scalable, maintainable reporting and analytics datasets.",
        "Create and maintain dbt models across silver and gold layers, improving data organization, reuse, and downstream reporting reliability.",
        "Optimize Snowflake SQL and dbt workloads, reducing data warehouse processing time by 20% during an initial improvement phase.",
        "Created a Kimball-based data warehouse modeling proof of concept to define standards and improve query and processing performance.",
        "Support MetricFlow semantic models to organize business metrics and maintain consistent analytical definitions.",
    ],
    "Snowflake, Snowflake SQL, dbt, MetricFlow, Kimball, Data Warehouse, SQL, Git",
)

role(
    doc,
    "Data Engineer",
    "Health Catalyst",
    "May 2021 - September 2024",
    "San Jose, Costa Rica",
    [
        "Prepared, cleaned, transformed, and modeled data for Power BI dashboards, business reporting, and operational decision-making.",
        "Designed customized ETL applications that reduced manual processing time by up to 40%.",
        "Built and maintained data workflows across SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI, and Excel.",
        "Automated exploratory analysis and validation routines, improving data-validation accuracy by 30% and strengthening report reliability.",
    ],
    "Power BI, SQL Server, T-SQL, SSIS, Python, Pandas, NumPy, Matplotlib, PowerShell, Snowflake, Excel",
)

role(
    doc,
    "SQL Developer",
    "Intertec International",
    "February 2019 - April 2021",
    "San Jose, Costa Rica",
    [
        "Developed and optimized complex SQL queries, stored procedures, and data workflows supporting analytics and reporting requirements.",
        "Consolidated disparate sources into analytics-ready datasets using optimized ETL processes.",
        "Improved data accuracy by more than 25% and reduced query execution time by up to 50% through validation, error handling, indexing, and SQL tuning.",
    ],
    "SQL Server, T-SQL, SSIS, Salesforce, MySQL",
)

role(
    doc,
    "DBA / SQL Developer",
    "EL Tiempo",
    "September 2020 - November 2020",
    "Colombia",
    [
        "Led migration activities for ten SQL Server databases while coordinating validation, data integrity, deployment timing, and operational continuity.",
    ],
    "SQL Server, SSIS, Azure, SSRS",
    page_break_before=True,
)

role(
    doc,
    "DBA / SQL and BI Consultant",
    "Gold Data Networks",
    "January 2016 - February 2020",
    "Panama City, Panama",
    [
        "Designed relational databases, reporting structures, stored procedures, and integration components for operational and BI requirements.",
        "Delivered SQL Server, SSIS, SSRS, Power BI, and data solutions aligned with stakeholder and infrastructure needs.",
    ],
    "Power BI, SQL Server, SSIS, SSRS, PostgreSQL, Excel",
)

role(
    doc,
    "Data Warehouse DBA",
    "BAC Credomatic",
    "November 2017 - January 2019",
    "San Jose, Costa Rica",
    [
        "Developed and maintained ETL pipelines that integrated multiple sources into the enterprise data warehouse.",
        "Designed and optimized tables, indexes, views, and analytical structures for reliable, high-performance reporting workloads.",
        "Analyzed large datasets to identify trends and patterns supporting business decision-making.",
    ],
    "Power BI, SQL Server, SSIS, SSAS, SSRS, Azure",
)

role(
    doc,
    "DBA and BI Consultant",
    "Bosal",
    "March 2017 - March 2018",
    "Lummen, Belgium",
    [
        "Designed the first stage of the company's primary data warehouse and developed reports and presentations for senior management.",
        "Worked with development teams to improve database practices and support BI and reporting delivery.",
    ],
    "Power BI, SQL Server, SSIS, Azure, MySQL",
)

section_heading(doc, "Additional Experience")
bullet(doc, "Xetux Solutions, Database Manager: established database policies and created a centralized sales data lake for reporting and analysis.")
bullet(doc, "ACH Cloud Services, SQL and BI Consultant: administered databases, created billing-management reports, and documented the data environment.")
bullet(doc, "EducaTablet and VIGEOSOFT, SQL Server DBA: administered production databases, modeled OLTP structures, improved response times, and trained development teams.")
bullet(doc, "Optica Caroni, Developer Analyst: implemented six operational projects and created management reports and dashboards.")

section_heading(doc, "Education and Languages")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
write(p, "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000 - 2004).")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
write(p, "Additional training: Analyzing and Visualizing Data with Power BI; Python for Data Analysis; Dataiku Core Designer; SQL Admin Part 1.")
p = doc.add_paragraph()
write(p, "Spanish: Native | English: Full Professional")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "Senior Power BI Developer CV"
doc.core_properties.subject = "Power BI, SQL, data modeling, ETL, reporting, and data warehouse development"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)

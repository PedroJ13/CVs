from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_Senior_Data_Analytics_SQL_Python_CV.docx"


def set_run_font(run, name="Calibri", size=9.3, color="1F2937", bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    set_run_font(r, size=10.4, color="000000", bold=True)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, size=9.1)


def add_role(doc, company, title, dates, location, bullets, technologies, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    r = p.add_run(f"{company} - {title}")
    set_run_font(r, size=10.0, color="111827", bold=True)
    r = p.add_run(f" | {dates} | {location}")
    set_run_font(r, size=9.0, color="4B5563")
    for item in bullets:
        add_bullet(doc, item)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Technologies: ")
    set_run_font(r, size=8.75, color="374151", bold=True)
    r = p.add_run(technologies)
    set_run_font(r, size=8.75, color="374151")


def add_skills_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    set_table_borders(table)
    set_cell_margins(table)
    for i, (label, value) in enumerate(rows):
        left, right = table.rows[i].cells
        set_cell_width(left, 2100)
        set_cell_width(right, 7260)
        shade(left, "E8EEF5")
        left.paragraphs[0].paragraph_format.space_after = Pt(0)
        right.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = left.paragraphs[0].add_run(label)
        set_run_font(r, size=8.65, color="000000", bold=True)
        r = right.paragraphs[0].add_run(value)
        set_run_font(r, size=8.65, color="1F2937")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(9.4)

for style_name in ("List Bullet", "List Paragraph"):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(9.1)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(1)
r = title.add_run("PEDRO JAVIER GUTIERREZ ARMAS")
set_run_font(r, size=16, color="000000", bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(1)
r = subtitle.add_run("Senior Data Analytics Specialist | SQL Server, Python, Power BI and Data Engineering")
set_run_font(r, size=10.2, color="000000", bold=True)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(5)
r = contact.add_run("San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13")
set_run_font(r, size=8.8, color="374151")

add_heading(doc, "Professional Summary")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.line_spacing = 1.05
summary = (
    "Senior Data Analytics Specialist, Data Engineer, and SQL Server professional with 15+ years of experience "
    "building data-informed reporting, BI, database, ETL/ELT, and data warehouse solutions. Strong hands-on background "
    "with SQL, SQL Server database administration and development, Python, Power BI, SSIS, SSAS, SSRS, Excel, "
    "PowerShell, Snowflake SQL, dbt, PostgreSQL, MySQL, Azure-based environments, data validation, statistical and "
    "exploratory analysis, and stakeholder-facing reporting. Experienced defining metrics, preparing and consolidating "
    "datasets, automating reporting workflows, improving data quality, interpreting trends and patterns, documenting "
    "processes, supporting users, and presenting actionable insights for operational and management decision-making."
)
r = p.add_run(summary)
set_run_font(r, size=9.15)

add_heading(doc, "Technical Skills")
add_skills_table(
    doc,
    [
        ("Analytics", "Descriptive analysis, trend identification, metrics definition, KPI reporting, data interpretation, actionable insights."),
        ("SQL Databases", "SQL Server, T-SQL, stored procedures, complex queries, views, indexes, query optimization, DBA support."),
        ("Python BI", "Python, Pandas, NumPy, Matplotlib, Power BI, SSRS, SSAS, Excel, dashboards, presentations."),
        ("Data Engineering", "ETL/ELT, SSIS, data extraction, refresh, cleansing, transformation, data warehouse pipelines, Snowflake SQL, dbt."),
        ("Delivery", "Stakeholder engagement, requirements translation, process documentation, user support, Git exposure, remote collaboration."),
    ],
)

add_heading(doc, "Professional Experience")
add_role(
    doc,
    "ServiceTitan",
    "Snowflake Developer / Data Engineer",
    "Sep 2024 - Present",
    "Remote / Costa Rica",
    [
        "Migrate reporting and business logic from C# processes into Snowflake SQL to support scalable reporting and analytics workflows.",
        "Create and maintain dbt models in silver and gold data layers, improving model organization, reusability, and downstream reporting reliability.",
        "Optimize Snowflake SQL and dbt workloads, reducing data warehouse processing time by approximately 20% in an initial optimization phase.",
    ],
    "Snowflake, Snowflake SQL, dbt, MetricFlow, Snowflake Cortex, SQL, C# logic migration, Cursor, AI-assisted PR review bot",
)
add_role(
    doc,
    "SMASH Costa Rica",
    "Data Engineer",
    "May 2021 - Sep 2024",
    "San Jose, Costa Rica",
    [
        "Designed and implemented customized ETL applications tailored to client requirements, reducing manual processing time by up to 40%.",
        "Built and maintained data workflows using SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI, and Excel.",
        "Developed automated exploratory data analysis and validation tools with Pandas, NumPy, and Matplotlib to detect anomalies and improve validation accuracy by 30%.",
        "Prepared, cleaned, transformed, and modeled datasets for Power BI dashboards, reporting, and operational decision-making.",
    ],
    "SQL Server, SSIS, Power BI, Python, Pandas, NumPy, Matplotlib, PowerShell, Snowflake, Excel",
)
add_role(
    doc,
    "Intertec International",
    "SQL Developer",
    "Feb 2019 - Apr 2021",
    "San Jose, Costa Rica",
    [
        "Developed and optimized SQL queries, stored procedures, and database workflows supporting analytics and business logic.",
        "Consolidated multiple disparate data sources into analytics-ready datasets using optimized ETL processes.",
        "Reduced query execution times by up to 50% through SQL optimization, indexing strategies, and performance tuning.",
    ],
    "SQL Server, T-SQL, SSIS, Salesforce, MySQL, stored procedures, query optimization, indexing",
)
add_role(
    doc,
    "BAC Credomatic",
    "Data Warehouse DBA",
    "Nov 2017 - Jan 2019",
    "San Jose, Costa Rica",
    [
        "Developed and maintained ETL pipelines to extract, transform, and load data from multiple sources into the data warehouse.",
        "Defined and optimized table, index, and view structures for high-performance analytical workloads.",
        "Improved operational efficiency by identifying trends and hidden patterns in large datasets for data-driven decision-making.",
    ],
    "SQL Server, SSIS, SSAS, Power BI, SSRS, Azure, data warehouse",
)
add_role(
    doc,
    "Gold Data Networks",
    "DBA / SQL & BI Consultant",
    "Jan 2016 - Feb 2020",
    "Panama City, Panama",
    [
        "Designed and implemented end-to-end data solutions integrated with existing infrastructure and BI/reporting needs.",
        "Built relational databases, table structures, custom database objects, and stored procedures for reporting and application workloads.",
        "Created Power BI and Excel reporting solutions using SQL Server, PostgreSQL, SSIS, and SSRS.",
    ],
    "SQL Server, T-SQL, SSIS, SSRS, PostgreSQL, Power BI, Excel",
)
add_role(
    doc,
    "EL Tiempo",
    "DBA / SQL Developer",
    "Sep 2020 - Nov 2020",
    "Colombia",
    [
        "Supported migration projects for 10 databases, coordinating with cross-functional teams to validate data and maintain integrity.",
        "Worked with SQL Server, SSIS, Azure, and SSRS to support reporting continuity and database migration activities.",
        "Documented database configurations and support practices to reduce operational risk during migration activities.",
    ],
    "SQL Server, SSIS, Azure, SSRS",
)

add_heading(doc, "Additional Relevant Experience")
for item in [
    "Earlier roles include DBA and BI Consultant at Bosal, Database Manager at Xetux Solutions, SQL and BI Consultant at ACH Cloud Services, SQL Server DBA roles at EducaTablet and VIGEOSOFT, and Developer Analyst at Optica Caroni C.A.",
    "Work included SQL Server administration and development, SSIS/SSRS assets, dashboards, management reports, billing reports, database modeling, documentation, process improvement, user support, and training development teams on database programming practices.",
    "Training includes Analyzing and Visualizing Data with Power BI, Python for Data Analysis, R Programming A-Z, Dataiku Core Designer, and SQL Admin Part 1.",
]:
    add_bullet(doc, item)

add_heading(doc, "Education and Languages")
for item in [
    "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio.",
    "English Certificate - Universidad Central de Venezuela / Microsoft.",
    "Spanish: Native | English: Full professional proficiency.",
]:
    add_bullet(doc, item)

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "Pedro Gutierrez - Senior Data Analytics SQL Python CV"
doc.save(OUTPUT)
print(OUTPUT)

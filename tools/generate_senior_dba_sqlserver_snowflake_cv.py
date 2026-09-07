from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from generate_fullstack_principal_sql_dba_cv import (
    OUTPUT_DIR,
    add_bullet,
    add_heading,
    add_role,
    add_run,
    configure_document,
    set_cell_borders,
    set_cell_margins,
    set_cell_shading,
    set_table_width,
)


DOCX_PATH = OUTPUT_DIR / "Pedro_Gutierrez_Senior_DBA_SQLServer_Snowflake_CV.docx"


def add_header(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("PEDRO JAVIER GUTIERREZ ARMAS")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(0)
    add_run(
        subtitle,
        "Senior Database Administrator | SQL Server | Snowflake | Data Warehousing",
        bold=True,
        size=10.5,
    )

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(5)
    add_run(
        contact,
        "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13",
        size=9,
    )


def add_skills_table(doc):
    rows = [
        (
            "Database Administration",
            "Microsoft SQL Server, database design, production support, maintenance, migrations, security practices, data integrity",
        ),
        (
            "SQL Performance",
            "T-SQL, complex queries, stored procedures, functions, views, indexing, query optimization, performance tuning, troubleshooting",
        ),
        (
            "Snowflake and Warehousing",
            "Snowflake SQL, warehouse workload optimization, dbt, MetricFlow, Kimball modeling, silver and gold layers, dimensional modeling",
        ),
        (
            "Data Engineering",
            "ETL/ELT pipelines, SSIS, SSAS, SSRS, data integration, data quality, validation, Power BI, Python, PowerShell",
        ),
        (
            "Cloud and AI Tools",
            "Azure-based environments, Snowflake Cortex, Cursor, AI-assisted PR review, Git-based collaborative workflows",
        ),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Area"
    header[1].text = "Tools and Practices"
    for cell in header:
        set_cell_shading(cell, "1F4E79")
        set_cell_borders(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(8.8)

    for area, detail in rows:
        cells = table.add_row().cells
        cells[0].text = area
        cells[1].text = detail
        for cell in cells:
            set_cell_borders(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
        cells[0].paragraphs[0].runs[0].font.bold = True
    set_table_width(table, [1.75, 5.65])


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_header(doc)

    add_heading(doc, "PROFESSIONAL SUMMARY")
    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(4)
    summary.add_run(
        "Senior Database Administrator, SQL Developer, and Data Engineer with 15+ years of experience designing, administering, optimizing, and supporting Microsoft SQL Server data solutions, production databases, ETL pipelines, and enterprise data warehouses. Hands-on experience with advanced T-SQL, stored procedures, indexing, query tuning, database migrations, data integrity, Snowflake SQL, dbt, and warehouse performance optimization. Works independently across remote and cross-functional teams, communicates technical findings clearly, and uses AI-assisted tools including Cursor, Snowflake Cortex, and automated PR review workflows to improve SQL development, documentation, and quality."
    )

    add_heading(doc, "TECHNICAL SKILLS")
    add_skills_table(doc)

    add_heading(doc, "PROFESSIONAL EXPERIENCE")
    add_role(
        doc,
        "Snowflake Developer / Data Engineer",
        "ServiceTitan",
        "September 2024 - Present",
        "Remote / Costa Rica",
        [
            "Develop and maintain Snowflake SQL and dbt models that transform reporting logic into scalable, reusable analytical assets.",
            "Optimize Snowflake and dbt workloads, reducing data warehouse processing time by 20% during an initial performance improvement phase.",
            "Created a Kimball-based data warehouse modeling proof of concept to improve structure, query behavior, reuse, and processing performance.",
            "Build and support silver and gold data layers, improving model organization, reliability, and downstream analytics consumption.",
            "Support MetricFlow semantic-layer development and Snowflake Cortex use cases for metric consistency and AI-enabled analytics workflows.",
        ],
        "Snowflake, Snowflake SQL, Snowflake Cortex, dbt, MetricFlow, Kimball, SQL, Cursor, AI-assisted PR review bot",
    )
    add_role(
        doc,
        "Data Engineer",
        "SMASH Costa Rica",
        "May 2021 - September 2024",
        "San Jose, Costa Rica",
        [
            "Built and maintained production data workflows across SQL Server, SSIS, Snowflake, Python, PowerShell, Power BI, and Excel.",
            "Designed ETL applications for client requirements, reducing manual data processing time by up to 40%.",
            "Prepared, transformed, modeled, and validated data for reporting, analytics, and operational decision-making.",
            "Developed automated anomaly detection and validation tools that improved data validation accuracy by 30%.",
        ],
        "SQL Server, SSIS, Snowflake, Python, PowerShell, Power BI, Pandas, NumPy, Excel",
    )
    add_role(
        doc,
        "SQL Developer",
        "Intertec International",
        "February 2019 - April 2021",
        "San Jose, Costa Rica",
        [
            "Developed and optimized advanced T-SQL queries, stored procedures, and database workflows supporting production business logic.",
            "Reduced query execution times by up to 50% through query tuning, indexing strategies, and performance analysis.",
            "Troubleshot database and data-integration issues and consolidated multiple source systems into analytics-ready datasets.",
            "Improved data accuracy by more than 25% through validation checks and error-handling mechanisms.",
        ],
        "SQL Server, T-SQL, SSIS, Salesforce, MySQL",
    )

    doc.add_page_break()
    add_role(
        doc,
        "DBA / SQL Developer",
        "EL Tiempo",
        "September 2020 - November 2020",
        "Colombia",
        [
            "Led planning and execution of migration activities for 10 SQL Server databases while maintaining data integrity and operational continuity.",
            "Coordinated cross-functional validation, deployment timing, risk identification, and production transition activities.",
            "Supported SQL Server and Azure-based database and reporting components across migration workflows.",
        ],
        "SQL Server, SSIS, Azure, SSRS",
    )
    add_role(
        doc,
        "DBA / SQL and BI Consultant",
        "Gold Data Networks",
        "January 2016 - February 2020",
        "Panama City, Panama",
        [
            "Designed relational databases, schemas, tables, stored procedures, and supporting objects for operational applications and reporting.",
            "Improved application performance through custom database logic and maintainable backend data structures.",
            "Delivered end-to-end database, integration, and BI solutions aligned with business and infrastructure requirements.",
        ],
        "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel",
    )
    add_role(
        doc,
        "Data Warehouse DBA",
        "BAC Credomatic",
        "November 2017 - January 2019",
        "San Jose, Costa Rica",
        [
            "Developed and maintained ETL pipelines that loaded multiple source systems into an enterprise data warehouse.",
            "Designed and optimized tables, indexes, and views for reliable, high-performance analytical workloads.",
            "Supported large datasets and reporting needs by improving structures used for trend analysis and business decision-making.",
        ],
        "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure",
    )

    add_heading(doc, "ADDITIONAL DATABASE EXPERIENCE")
    for item in [
        "Bosal, DBA and BI Consultant: designed the first stage of the main data warehouse, produced senior-management reports, and supported SQL Server, SSIS, Azure, MySQL, and Power BI solutions.",
        "Xetux Solutions, Database Manager: defined database management policies, improved performance and data security, and created a data lake for centralized sales reporting.",
        "ACH Cloud Services, SQL and BI Consultant: administered and developed SQL Server databases, created billing dashboards, and documented the data environment.",
        "EducaTablet, Microsoft SQL Server DBA: administered production databases, improved database process response times, and trained developers in database programming practices.",
        "VIGEOSOFT, Microsoft SQL Server DBA: modeled, designed, configured, and programmed OLTP databases for business applications.",
        "Optica Caroni, Developer Analyst: delivered SQL Server, SSIS, SSRS, .NET, and VB6 solutions supporting production workflows and management reporting.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "EDUCATION AND CERTIFICATIONS")
    for item in [
        "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000 - 2004)",
        "English Certificate - Universidad Central de Venezuela / Microsoft (2013 - 2015)",
        "Analyzing and Visualizing Data with Power BI; Complete Data Science Training with Python for Data Analysis; Dataiku Core Designer; SQL Admin Part 1",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "LANGUAGES")
    doc.add_paragraph("Spanish: Native | English: Full Professional")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()

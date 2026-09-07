from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT_DIR = Path(r"C:\Work\CVs\Output")
DOCX_PATH = OUTPUT_DIR / "Pedro_Gutierrez_FullStack_Principal_SQL_DBA_CV.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_run(paragraph, text, bold=False, italic=False, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    return run


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.style = "Heading 1"
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.02
    p.add_run(text)
    return p


def add_role(doc, title, company, dates, location, bullets, tech=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, f"{title} | {company}", bold=True)
    add_run(p, f"\n{dates} | {location}", italic=True, size=9)
    for bullet in bullets:
        add_bullet(doc, bullet)
    if tech:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_run(p, "Technologies: ", bold=True, size=9)
        add_run(p, tech, size=9)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(9.2)
    styles["Normal"].font.color.rgb = None
    styles["Normal"].paragraph_format.space_after = Pt(2)
    styles["Normal"].paragraph_format.line_spacing = 1.03

    styles["Title"].font.name = "Calibri"
    styles["Title"].font.size = Pt(17)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = None
    styles["Title"].paragraph_format.space_after = Pt(0)

    styles["Heading 1"].font.name = "Calibri"
    styles["Heading 1"].font.size = Pt(10.5)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = None
    styles["Heading 1"].paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2"):
        styles[style_name].font.name = "Calibri"
        styles[style_name].font.size = Pt(9)
        styles[style_name].paragraph_format.left_indent = Inches(0.18)
        styles[style_name].paragraph_format.first_line_indent = Inches(-0.12)


def add_header(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("PEDRO JAVIER GUTIERREZ ARMAS")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(0)
    add_run(
        subtitle,
        "Principal SQL Database Administrator | SQL Server DBA | T-SQL Performance Tuning",
        bold=True,
        size=10.5,
    )

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(5)
    add_run(
        contact,
        "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13",
        size=9,
    )


def add_skills_table(doc):
    rows = [
        (
            "SQL Server Administration",
            "SQL Server, database health management, maintenance routines, production support, migration support, multi-database environments, Azure-based environments",
        ),
        (
            "Performance Engineering",
            "T-SQL optimization, stored procedure tuning, indexing strategies, execution plan review, slow-query troubleshooting, bottleneck analysis",
        ),
        (
            "Database Development",
            "Stored procedures, functions, views, schema design, relational modeling, data validation, data integrity, SQL-based business logic",
        ),
        (
            "Data Engineering and BI",
            "SSIS, SSAS, SSRS, ETL/ELT pipelines, data warehousing, Power BI, Snowflake SQL, dbt, PostgreSQL, MySQL",
        ),
        (
            "Automation and Delivery",
            "Python, PowerShell, Git exposure, Agile/Scrum collaboration, documentation, Cursor, AI-assisted PR review bots",
        ),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Area"
    hdr[1].text = "Tools and Practices"
    for cell in hdr:
        set_cell_shading(cell, "1F4E79")
        set_cell_borders(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = None
                run.font.size = Pt(8.8)
    for area, detail in rows:
        cells = table.add_row().cells
        cells[0].text = area
        cells[1].text = detail
        for cell in cells:
            set_cell_borders(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
        cells[0].paragraphs[0].runs[0].font.bold = True
    set_table_width(table, [1.75, 5.65])


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_header(doc)

    add_heading(doc, "PROFESSIONAL SUMMARY")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "Principal SQL Server DBA, SQL Developer, and Data Engineer with 15+ years of experience designing, administering, optimizing, and maintaining data platforms for reporting, operational systems, ETL/ELT, data warehousing, and BI. Strong background in Microsoft SQL Server, advanced T-SQL, stored procedures, indexing, query optimization, database maintenance, production troubleshooting, data migrations, and stakeholder-facing delivery. Experienced working in remote Agile environments with autonomy, supporting complex SQL Server and Azure-based data environments, defining database practices, documenting technical processes, and using automation and AI-assisted development tools such as Cursor and PR review bots."
    )

    add_heading(doc, "TECHNICAL SKILLS")
    add_skills_table(doc)

    add_heading(doc, "PROFESSIONAL EXPERIENCE")
    add_role(
        doc,
        "Snowflake Developer / Data Engineer",
        "ServiceTitan",
        "September 2024 - Present",
        "Remote / Costa Rica",
        [
            "Migrated C# reporting logic into Snowflake SQL, improving maintainability of reporting logic and supporting scalable analytics workflows.",
            "Create and maintain dbt models in silver and gold layers, improving model organization, reuse, and downstream reporting reliability.",
            "Optimized Snowflake SQL and dbt workloads, reducing data warehouse processing time by 20% during an initial performance improvement phase.",
            "Created a Kimball-based data warehouse model proof of concept to improve modeling standards, query behavior, and processing performance.",
            "Support MetricFlow semantic layer development and Snowflake Cortex use cases for AI-enabled analytics, metric consistency, and data documentation.",
        ],
        "Snowflake, Snowflake SQL, Snowflake Cortex, dbt, MetricFlow, Kimball, SQL, C# logic migration, Cursor, AI-assisted PR review bot",
    )
    add_role(
        doc,
        "Data Engineer",
        "SMASH Costa Rica",
        "May 2021 - September 2024",
        "San Jose, Costa Rica",
        [
            "Designed and implemented ETL applications tailored to client requirements, reducing manual data processing time by up to 40%.",
            "Built and maintained production data workflows using SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI, and Excel.",
            "Developed automated exploratory analysis and validation tools with Pandas, NumPy, and Matplotlib to detect anomalies and improve validation accuracy by 30%.",
            "Prepared, cleaned, transformed, and modeled datasets for business reporting, dashboards, and operational decision-making.",
        ],
        "SQL Server, SSIS, Power BI, Python, Pandas, NumPy, Matplotlib, PowerShell, Snowflake, Excel",
    )
    add_role(
        doc,
        "SQL Developer",
        "Intertec International",
        "February 2019 - April 2021",
        "San Jose, Costa Rica",
        [
            "Developed and optimized stored procedures, complex SQL queries, and database workflows supporting business logic and analytics workloads.",
            "Reduced query execution times by up to 50% through T-SQL tuning, indexing strategies, and database performance optimization.",
            "Consolidated disparate data sources into analytics-ready datasets using optimized ETL processes.",
            "Improved data accuracy by more than 25% through validation checks and error-handling mechanisms across key workflows.",
        ],
        "SQL Server, T-SQL, SSIS, Salesforce, MySQL",
    )
    doc.add_page_break()
    add_role(
        doc,
        "DBA / SQL Developer",
        "EL Tiempo",
        "September 2020 - November 2020",
        "Colombia",
        [
            "Led planning and execution of migration activities for 10 SQL Server databases, coordinating validation, deployment timing, and continuity of operations.",
            "Worked with cross-functional teams to identify risks, validate migrated data, and preserve database integrity during production changes.",
            "Supported SQL Server and Azure-based database/reporting environments with SSIS and SSRS components.",
        ],
        "SQL Server, SSIS, Azure, SSRS",
    )
    add_role(
        doc,
        "DBA / SQL and BI Consultant",
        "Gold Data Networks",
        "January 2016 - February 2020",
        "Panama City, Panama",
        [
            "Built relational databases and table structures from the ground up to support dynamic web applications and reporting workloads.",
            "Developed custom database objects and stored procedures to improve application performance and support scalable backend processes.",
            "Designed end-to-end data solutions integrated with existing infrastructure, BI needs, and operational reporting.",
        ],
        "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel",
    )
    add_role(
        doc,
        "Data Warehouse DBA",
        "BAC Credomatic",
        "November 2017 - January 2019",
        "San Jose, Costa Rica",
        [
            "Developed and maintained ETL pipelines to load data from multiple sources into the enterprise data warehouse.",
            "Defined and optimized table, index, and view structures for high-performance analytical workloads.",
            "Supported reporting and analysis needs by identifying trends and patterns in large datasets for business decision-making.",
        ],
        "SQL Server, SSIS, SSAS, Power BI, SSRS, Azure",
    )

    add_heading(doc, "ADDITIONAL SQL SERVER DBA AND DATABASE EXPERIENCE")
    for text in [
        "Bosal, DBA and BI Consultant: designed the first stage of the main data warehouse, generated senior-management reports, and shared database improvement practices with development teams. Technologies: SQL Server, SSIS, Azure, MySQL, Power BI.",
        "Xetux Solutions, Database Manager: defined database management policies and improved performance and data security across business areas. Technologies: SQL Server, SSIS, SSRS, Azure.",
        "ACH Cloud Services, SQL and BI Consultant: administered and developed company databases, created billing dashboards, and documented the data environment. Technologies: SQL Server, SSIS, SSRS, Azure.",
        "EducaTablet, Microsoft SQL Server DBA: administered and developed company databases, improved response times for database processes, and trained development teams on database programming practices.",
        "VIGEOSOFT, Microsoft SQL Server DBA: modeled, designed, configured, and programmed OLTP databases for business applications, including data dictionary practices and stakeholder coordination.",
        "Optica Caroni C.A., Developer Analyst: implemented production workflow and reporting projects using SQL Server, SSIS, SSRS, .NET, and VB6.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "EDUCATION AND CERTIFICATIONS")
    for text in [
        "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000 - 2004)",
        "English Certificate - Universidad Central de Venezuela / Microsoft (2013 - 2015)",
        "Analyzing and Visualizing Data with Power BI; Complete Data Science Training with Python for Data Analysis; R Programming A-Z; Dataiku Core Designer; SQL Admin Part 1",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "LANGUAGES")
    doc.add_paragraph("Spanish: Native | English: Full Professional")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()

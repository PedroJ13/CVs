from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_EXSquared_Senior_DBA_SQL_Developer_CV.docx"


def set_run_font(run, name="Calibri", size=9.3, color="1F2937", bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    set_run_font(r, size=10.4, color="000000", bold=True)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(2.1)
    p.paragraph_format.line_spacing = 1.03
    r = p.add_run(text)
    set_run_font(r, size=9.1)


def add_role(doc, company, title, dates, location, bullets, technologies, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    r = p.add_run(f"{company} - {title}")
    set_run_font(r, size=10.0, color="111827", bold=True)
    r = p.add_run(f" | {dates} | {location}")
    set_run_font(r, size=9.0, color="4B5563")
    for item in bullets:
        add_bullet(doc, item)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Technologies: ")
    set_run_font(r, size=8.75, color="374151", bold=True)
    r = p.add_run(technologies)
    set_run_font(r, size=8.75, color="374151")


def add_skills_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    set_table_borders(table)
    set_cell_margins(table)
    for i, (label, value) in enumerate(rows):
        left, right = table.rows[i].cells
        set_cell_width(left, 2100)
        set_cell_width(right, 7260)
        shade(left, "E8EEF5")
        left.paragraphs[0].paragraph_format.space_after = Pt(0)
        right.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = left.paragraphs[0].add_run(label)
        set_run_font(r, size=8.65, color="000000", bold=True)
        r = right.paragraphs[0].add_run(value)
        set_run_font(r, size=8.65, color="1F2937")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(9.4)

for style_name in ("List Bullet", "List Paragraph"):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(9.1)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(1)
r = title.add_run("PEDRO JAVIER GUTIERREZ ARMAS")
set_run_font(r, size=16, color="000000", bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(1)
r = subtitle.add_run("Senior Database Administrator and SQL Developer")
set_run_font(r, size=10.4, color="000000", bold=True)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(5)
r = contact.add_run("San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13")
set_run_font(r, size=8.8, color="374151")

add_heading(doc, "Professional Summary")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.line_spacing = 1.05
summary = (
    "Senior SQL Server DBA, SQL Developer, and Data Engineer with 15+ years of experience designing, developing, "
    "optimizing, troubleshooting, and maintaining database solutions for business-critical reporting, BI, ETL, and "
    "data warehouse environments. Strong hands-on background with Microsoft SQL Server, Azure SQL Database and "
    "Azure-based database environments, advanced T-SQL, stored procedures, functions, views, schema design, indexing, "
    "query optimization, execution behavior analysis, SSIS, SSRS, SSAS, Power BI, Python, PowerShell, PostgreSQL, "
    "MySQL, Snowflake SQL, and large analytical datasets. Experienced resolving production database issues, improving "
    "performance, maintaining data integrity, transforming and moving data across systems, and collaborating with "
    "technical and business stakeholders in advanced English."
)
r = p.add_run(summary)
set_run_font(r, size=9.2)

add_heading(doc, "Technical Skills")
add_skills_table(
    doc,
    [
        ("SQL Server", "Microsoft SQL Server, Azure SQL Database, Azure-based database environments, database administration, database development."),
        ("T-SQL Development", "Complex queries, stored procedures, functions, views, database business logic, schema evolution, database objects."),
        ("Performance", "Query optimization, indexing strategies, execution behavior review, performance tuning, response-time improvement."),
        ("Troubleshooting", "Production support, slow-query analysis, database workflow issues, validation checks, error handling, data integrity."),
        ("Data Movement", "SSIS, ETL/ELT, data transformation, data migration, multi-source consolidation, reporting-ready datasets."),
        ("BI and Automation", "SSRS, SSAS, Power BI, Excel, Python, PowerShell, Snowflake SQL, PostgreSQL, MySQL, Salesforce integrations."),
    ],
)

add_heading(doc, "Professional Experience")
add_role(
    doc,
    "ServiceTitan",
    "Snowflake Developer / Data Engineer",
    "Sep 2024 - Present",
    "Remote / Costa Rica",
    [
        "Migrate reporting and business logic from C# processes into Snowflake SQL, improving maintainability and traceability of analytics workflows.",
        "Create and maintain dbt models in silver and gold data layers to support reusable, reliable, and analytics-ready datasets.",
        "Optimize Snowflake SQL and dbt workloads, reducing data warehouse processing time by approximately 20% in an initial optimization phase.",
        "Support MetricFlow semantic layer work with Snowflake Cortex, improving metric organization and consistency for analytics consumers.",
    ],
    "Snowflake, Snowflake SQL, dbt, MetricFlow, Snowflake Cortex, SQL, C# logic migration, Cursor, AI-assisted PR review bot",
)
add_role(
    doc,
    "SMASH Costa Rica",
    "Data Engineer",
    "May 2021 - Sep 2024",
    "San Jose, Costa Rica",
    [
        "Designed and maintained ETL workflows using SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI, and Excel, reducing manual processing time by up to 40%.",
        "Built automated data refresh, validation, and exploratory analysis workflows to detect anomalies and improve data validation accuracy by 30%.",
        "Prepared, cleaned, transformed, and modeled datasets for business reporting, Power BI dashboards, and operational decision-making.",
        "Collaborated with stakeholders and technical teams to translate reporting needs into repeatable SQL and data routines.",
    ],
    "SQL Server, SSIS, T-SQL, Python, PowerShell, Snowflake, Power BI, Excel, Pandas, NumPy, Matplotlib",
)
add_role(
    doc,
    "Intertec International",
    "SQL Developer",
    "Feb 2019 - Apr 2021",
    "San Jose, Costa Rica",
    [
        "Developed and optimized SQL Server stored procedures, complex T-SQL queries, and database workflows supporting business logic and analytics needs.",
        "Reduced query execution times by up to 50% through SQL optimization, indexing strategies, and performance tuning.",
        "Analyzed and resolved data workflow issues using validation checks, error handling, and structured troubleshooting across key processes.",
        "Integrated disparate data sources, including Salesforce and MySQL, into analytics-ready datasets using optimized ETL processes.",
    ],
    "SQL Server, T-SQL, SSIS, stored procedures, Salesforce, MySQL, query optimization, indexing",
)
add_role(
    doc,
    "EL Tiempo",
    "DBA / SQL Developer",
    "Sep 2020 - Nov 2020",
    "Colombia",
    [
        "Supported SQL Server database migration and upgrade activities for 10 databases, including assessment, validation, and cutover support.",
        "Worked with SQL Server and Azure SQL environments, helping with configuration, maintenance, monitoring, and operational continuity.",
        "Validated data integrity and coordinated with technical teams to reduce operational risk during migration activities.",
    ],
    "SQL Server, SSIS, Azure, SSRS",
    page_break_before=True,
)
add_role(
    doc,
    "Gold Data Networks",
    "DBA / SQL & BI Consultant",
    "Jan 2016 - Feb 2020",
    "Panama City, Panama",
    [
        "Built relational databases, table structures, custom database objects, and stored procedures from the ground up for application and reporting workloads.",
        "Delivered SQL Server, SSIS, SSRS, PostgreSQL, Power BI, and Excel solutions integrated with existing infrastructure and BI needs.",
        "Tuned SQL queries, indexes, and database objects to improve application and reporting performance.",
    ],
    "SQL Server, T-SQL, SSIS, SSRS, PostgreSQL, Power BI, Excel",
)
add_role(
    doc,
    "BAC Credomatic",
    "Data Warehouse DBA",
    "Nov 2017 - Jan 2019",
    "San Jose, Costa Rica",
    [
        "Developed and maintained ETL pipelines to extract, transform, and load data from multiple sources into the data warehouse.",
        "Defined and optimized table, index, and view structures for high-performance analytical workloads.",
        "Supported SSAS, Power BI, SSRS, and Azure-based reporting environments for large-scale analytics and decision support.",
    ],
    "SQL Server, SSIS, SSAS, Power BI, SSRS, Azure, data warehouse",
)

add_heading(doc, "Additional Database Experience")
for item in [
    "Administered and developed SQL Server databases, reporting workflows, SSIS/SSRS assets, database policies, data dictionaries, and production process improvements across Xetux Solutions, ACH Cloud Services, EducaTablet, VIGEOSOFT, and Optica Caroni C.A.",
    "Modeled, designed, configured, and programmed OLTP databases for business applications and trained development teams on database programming good practices.",
    "Created a data lake to consolidate sales data and support centralized reporting and analysis.",
]:
    add_bullet(doc, item)

add_heading(doc, "Education Certifications and Languages")
for item in [
    "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio.",
    "English Certificate - Universidad Central de Venezuela / Microsoft.",
    "SQL Admin Part 1; Analyzing and Visualizing Data with Power BI; Complete Data Science Training with Python for Data Analysis; R Programming A-Z; Dataiku Core Designer.",
    "Spanish: Native | English: Full professional proficiency.",
]:
    add_bullet(doc, item)

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "Pedro Gutierrez - EX Squared Senior DBA SQL Developer CV"
doc.save(OUTPUT)
print(OUTPUT)

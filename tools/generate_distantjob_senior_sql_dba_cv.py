from docx import Document
from docx.shared import Pt

import generate_senior_microsoft_sql_server_dba_partner_cv as source


DOCX_PATH = source.OUTPUT_DIR / "Pedro_Gutierrez_DistantJob_Senior_SQL_DBA_CV.docx"


def replace_paragraph(paragraph, text, bold=False):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    run.bold = bold


def build_document():
    source.DOCX_PATH = DOCX_PATH
    source.build_document()

    doc = Document(DOCX_PATH)
    for paragraph in doc.paragraphs:
        if paragraph.text == "Senior Microsoft SQL Server DBA | T-SQL Performance | Data Platforms":
            replace_paragraph(
                paragraph,
                "Senior SQL Server DBA | Production Support | T-SQL Performance",
                bold=True,
            )
        elif paragraph.text.startswith("Senior Microsoft SQL Server DBA, SQL Developer, and Data Engineer"):
            replace_paragraph(
                paragraph,
                "Senior SQL Server DBA and SQL Developer with 15+ years of experience supporting production databases, developing database solutions, optimizing T-SQL workloads, and maintaining data platforms for operational applications, ETL, reporting, and analytics. Strong hands-on background in stored procedures, indexing, slow-query tuning, database migrations, Always On, Extended Events, Query Store, SSIS, PowerShell, data integrity, security practices, and Azure-based environments. Experienced partnering with developers to improve database logic and performance, coordinating production changes, documenting technical processes, reviewing AI-assisted pull requests, and training development teams on database programming practices.",
            )

    replacements = {
        "SQL Server Administration": (
            "SQL Server Operations",
            "Microsoft SQL Server, production support, database maintenance, multi-database migrations, Always On, operational continuity",
        ),
        "Performance Engineering": (
            "Performance and Monitoring",
            "Advanced T-SQL, stored procedures, indexing, slow-query tuning, workload analysis, Extended Events, Query Store, troubleshooting",
        ),
        "Data Architecture and BI": (
            "Database Development",
            "Relational and OLTP modeling, schemas, tables, views, database objects, SSIS, ETL/ELT, data warehousing",
        ),
        "Automation and Cloud": (
            "Automation and Platforms",
            "PowerShell, Python, Azure-based environments, SQL Server, Snowflake, SSIS, Git-based collaboration",
        ),
        "Data Operations": (
            "Delivery and Collaboration",
            "Production issue resolution, data integrity, validation, security practices, technical documentation, developer guidance, stakeholder coordination",
        ),
    }

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            area = row.cells[0].text
            if area in replacements:
                new_area, details = replacements[area]
                row.cells[0].text = new_area
                row.cells[1].text = details
            if area != "Area":
                for run in row.cells[0].paragraphs[0].runs:
                    run.bold = True
                for cell in row.cells:
                    for cell_paragraph in cell.paragraphs:
                        for run in cell_paragraph.runs:
                            run.font.size = Pt(8.5)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()

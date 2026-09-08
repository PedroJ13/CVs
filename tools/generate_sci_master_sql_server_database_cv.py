from docx import Document

import generate_senior_microsoft_sql_server_dba_partner_cv as source


DOCX_PATH = source.OUTPUT_DIR / "Pedro_Gutierrez_SCI_Master_SQL_Server_Database_CV.docx"


def replace_paragraph(paragraph, text, bold=False):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    run.bold = bold


def build_document():
    source.DOCX_PATH = DOCX_PATH
    source.build_document()

    doc = Document(DOCX_PATH)
    for paragraph in doc.paragraphs:
        if paragraph.text == "Senior Microsoft SQL Server DBA | T-SQL Performance | Data Platforms":
            replace_paragraph(
                paragraph,
                "Master SQL Server Database Developer and Administrator",
                bold=True,
            )
        elif paragraph.text.startswith("Senior Microsoft SQL Server DBA, SQL Developer, and Data Engineer"):
            replace_paragraph(
                paragraph,
                "Master-level SQL Server Database Developer and Administrator with 15+ years of experience designing, developing, optimizing, and supporting database solutions for production applications, ETL, data warehousing, and business reporting. Strong hands-on background in Microsoft SQL Server, advanced T-SQL, stored procedures, relational and OLTP modeling, indexing, query tuning, Always On, Extended Events, Query Store, SSIS, PowerShell, database migrations, data integrity, security practices, and Azure-based environments. Proven ability to translate business requirements into reliable database solutions, troubleshoot performance and data issues, coordinate releases and migrations, document technical standards, and guide development teams on database programming and optimization practices.",
            )

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()

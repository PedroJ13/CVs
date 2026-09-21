from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_Delphos_DBA_SQL_Server_CV_EN.docx"


def text(paragraph, value, size=9.2, bold=False, color="25313D"):
    part = paragraph.add_run(value)
    part.font.name = "Calibri"
    part.font.size = Pt(size)
    part.font.bold = bold
    part.font.color.rgb = RGBColor.from_string(color)


def heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    text(p, title.upper(), 10.5, True, "174B70")


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.6)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    text(p, value)


def role(doc, title, company, dates, location, points, technologies, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4.5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    text(p, f"{title} | {company}", 10, True, "17365D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    text(p, f"{dates} | {location}", 8.7, color="596878")
    for point in points:
        bullet(doc, point)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    text(p, "Technologies: ", 8.6, True, "596878")
    text(p, technologies, 8.6, color="596878")


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.52)
sec.bottom_margin = Inches(0.52)
sec.left_margin = Inches(0.65)
sec.right_margin = Inches(0.65)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
text(p, "PEDRO JAVIER GUTIERREZ ARMAS", 16, True, "17365D")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
text(p, "Database Administrator | SQL Server and SQL Development", 10.1, True, "174B70")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
text(p, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.6, color="596878")

heading(doc, "Professional Summary")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.04
text(p, "Database professional with more than 15 years of experience administering and developing SQL Server solutions in enterprise environments. Skilled in relational database design, SQL and T-SQL queries, joins, indexes, stored procedures, ETL, performance optimization, migrations, data integrity, and technical documentation. Experienced collaborating with development and business teams to resolve data issues and support operational and reporting processes. Based in Costa Rica with remote-work experience.")

heading(doc, "Technical Skills")
for label, value in [
    ("Databases", "SQL Server, PostgreSQL, MySQL, Snowflake; table, view, and index design."),
    ("SQL development", "T-SQL, complex queries, joins, stored procedures, and query optimization."),
    ("Administration", "Migrations, data validation and integrity, performance analysis, process documentation, and data dictionaries."),
    ("Integration", "SSIS, SSRS, SSAS, ETL, Power BI, Python, and PowerShell."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    text(p, f"{label}: ", 9.1, True)
    text(p, value, 9.1)

heading(doc, "Professional Experience")
role(doc, "Snowflake Developer / Data Engineer", "ServiceTitan", "Sep 2024 - Present", "Remote / Costa Rica", [
    "Migrate reporting logic from C# to Snowflake SQL and maintain dbt models for reusable analytical data.",
    "Optimize queries and warehouse processing; achieved an initial 20% reduction in processing time.",
], "Snowflake SQL, dbt, data modeling, Kimball")
role(doc, "Data Engineer", "SMASH Costa Rica", "May 2021 - Sep 2024", "San Jose, Costa Rica", [
    "Designed and maintained ETL applications and data workflows using SQL Server and SSIS to meet client requirements.",
    "Automated data validation and analysis, improving validation accuracy by 30% and reducing manual processing by up to 40%.",
], "SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI")
role(doc, "SQL Developer", "Intertec International", "Feb 2019 - Apr 2021", "San Jose, Costa Rica", [
    "Developed and optimized stored procedures, SQL queries, and database workflows for business logic and reporting.",
    "Reduced query execution times by up to 50% through performance analysis, SQL tuning, and indexing strategies.",
], "SQL Server, T-SQL, SSIS, Salesforce, MySQL")
role(doc, "DBA / SQL Developer", "EL Tiempo", "Sep 2020 - Nov 2020", "Colombia", [
    "Led the migration of ten databases, validating data, risks, and integrity during implementation.",
], "SQL Server, SSIS, Azure, SSRS", page_break_before=True)
role(doc, "DBA / SQL & BI Consultant", "Gold Data Networks", "Jan 2016 - Feb 2020", "Panama City, Panama", [
    "Built relational databases, tables, and SQL objects for applications and reporting workloads; developed stored procedures.",
], "SQL Server, SSIS, SSRS, PostgreSQL, Power BI")
role(doc, "Data Warehouse DBA", "BAC Credomatic", "Nov 2017 - Jan 2019", "San Jose, Costa Rica", [
    "Maintained ETL processes and defined tables, indexes, and views to improve analytical workload performance.",
], "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure")

heading(doc, "Additional Experience")
bullet(doc, "EducaTablet: administered and developed SQL Server databases and trained developers in database programming practices.")
bullet(doc, "VIGEOSOFT: modeled, designed, and configured OLTP databases; documented structural changes using data dictionaries.")
bullet(doc, "ACH Cloud Services, Bosal, Xetux Solutions, and Optica Caroni: database administration, data integration, and enterprise reporting.")

heading(doc, "Education and Languages")
p = doc.add_paragraph()
text(p, "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000-2004).")
p = doc.add_paragraph()
text(p, "Additional training: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI.")
p = doc.add_paragraph()
text(p, "Spanish: Native | English: Full professional proficiency.")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "Delphos SQL Server DBA CV English"
doc.save(OUTPUT)
print(OUTPUT)

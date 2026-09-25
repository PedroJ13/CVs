from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Work\CVs")
SOURCE = ROOT / "Base" / "Pedro_Gutierrez_CV_SQL_Server_DBA.docx"
OUTPUT = ROOT / "Output" / "Pedro_Gutierrez_Insight_Global_SSIS_Integration_Engineer_CV.docx"

BLUE = "1F4E79"
DARK = "222222"
GRAY = "5B5B5B"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def write(paragraph, text, size=9.1, bold=False, italic=False, color=None, font="Aptos"):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_bottom_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "BDD7EE")
    borders.append(bottom)


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    write(p, text.upper(), 10.2, True, color=BLUE, font="Aptos Display")
    add_bottom_rule(p)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.4)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, text)


def skill_line(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.6)
    write(p, f"{label}: ", 9.0, True)
    write(p, text, 9.0)


def role(doc, title, company, dates, place, points, technologies, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    write(p, f"{title} | {company}", 10, True, color=DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    write(p, f"{dates} | {place}", 8.5, italic=True, color=GRAY)

    for point in points:
        bullet(doc, point)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write(p, "Technologies: ", 8.6, True, color=DARK)
    write(p, technologies, 8.6, color=GRAY)


doc = Document(SOURCE)
clear_body(doc)
page = doc.sections[0]
page.top_margin = Inches(0.48)
page.bottom_margin = Inches(0.48)
page.left_margin = Inches(0.62)
page.right_margin = Inches(0.62)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9.1)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0.5)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 17, True, color=BLUE, font="Aptos Display")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "SSIS Integration Engineer | SQL Server, ETL and Enterprise Data Pipelines", 10, True, color="404040")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
write(p, "San Jose, Costa Rica (CST) | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.5, color="4C4C4C")

section_heading(doc, "Professional Summary")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.05
p.paragraph_format.space_after = Pt(4)
write(
    p,
    "SSIS integration engineer, SQL Server specialist, and data professional with 15+ years of experience building, supporting, testing, and optimizing enterprise data integrations. Strong hands-on background in SQL Server Integration Services, advanced T-SQL, ETL/ELT pipelines, multi-source data integration, stored procedures, production support, data validation, performance tuning, migrations, automation, and technical documentation. Proven record of reducing manual processing, improving data quality, troubleshooting live integration issues, and coordinating controlled deployments with developers, QA, support, infrastructure, and business stakeholders. Based in Costa Rica and aligned with CST business hours.",
)

section_heading(doc, "Technical Skills")
skill_line(doc, "SSIS and integration", "SSIS package development and support, ETL/ELT, data pipelines, multi-source integration, scheduling, validation, error handling, deployment support, and production troubleshooting")
skill_line(doc, "SQL Server", "advanced T-SQL, stored procedures, functions, views, indexing, query optimization, execution analysis, SQL Agent, backup and restore, and database administration")
skill_line(doc, "Quality and testing", "data reconciliation, anomaly detection, source-to-target validation, regression support, exception diagnostics, post-deployment verification, and operational monitoring")
skill_line(doc, "Automation and platforms", "PowerShell, Python, Git, AWS RDS, Azure-based environments, Snowflake, MySQL, PostgreSQL, Power BI, SSAS, and SSRS")
skill_line(doc, "Delivery", "requirements translation, runbooks, technical documentation, stakeholder communication, controlled releases, incident response, and cross-functional collaboration")

section_heading(doc, "Professional Experience")
role(
    doc,
    "SQL Server DBA / Database Engineer",
    "MWR Life",
    "September 2024 - Present",
    "Remote / Costa Rica",
    [
        "Develop and support T-SQL stored procedures, functions, deployment scripts, diagnostic utilities, and database processes used by production applications and integrations.",
        "Led an InEvent data integration remediation effort covering source backup, data validation, controlled deployment, recovery of hundreds of records, stored procedure corrections, and post-release verification.",
        "Built diagnostic tooling to identify missing records, failed integrations, duplicate data, invalid dates, unmatched entities, and webhook-related data-quality exceptions.",
        "Monitor SQL Agent jobs, scheduled processes, long-running queries, execution history, alerts, and production incidents across SQL Server environments hosted on AWS RDS.",
    ],
    "SQL Server, AWS RDS, T-SQL, SQL Agent, PowerShell, JSON/OpenJSON, Git, monitoring, data validation, production support",
)

role(
    doc,
    "DBA / SQL Developer",
    "Health Catalyst",
    "May 2021 - September 2024",
    "San Jose, Costa Rica",
    [
        "Developed and supported production data workflows and client-specific integrations using SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI, and Excel.",
        "Designed customized ETL applications that reduced manual processing time by up to 40%.",
        "Automated data validation and anomaly-detection routines, improving validation accuracy by 30% and strengthening reporting reliability.",
        "Prepared, transformed, and modeled data from multiple sources for operational reporting and downstream analytics.",
    ],
    "SQL Server, SSIS, T-SQL, Python, PowerShell, Snowflake, Power BI, Pandas, Excel",
)

role(
    doc,
    "SQL Developer",
    "Intertec International",
    "February 2019 - April 2021",
    "San Jose, Costa Rica",
    [
        "Developed and optimized T-SQL queries, stored procedures, SSIS workflows, and integrations across SQL Server, Salesforce, and MySQL.",
        "Consolidated disparate sources into analytics-ready datasets and improved data accuracy by more than 25% through validation and error handling.",
        "Reduced query execution times by up to 50% through indexing, query optimization, and performance tuning.",
    ],
    "SQL Server, T-SQL, SSIS, Salesforce, MySQL",
)

role(
    doc,
    "DBA / SQL Developer",
    "EL Tiempo",
    "September 2020 - November 2020",
    "Colombia",
    [
        "Led planning and execution of migration activities for ten SQL Server databases while maintaining data integrity and operational continuity.",
        "Coordinated migration risks, validation, deployment timing, and transition activities with cross-functional teams.",
    ],
    "SQL Server, SSIS, Azure, SSRS",
    page_break_before=True,
)

role(
    doc,
    "DBA / SQL and BI Consultant",
    "Gold Data Networks",
    "January 2016 - February 2020",
    "Panama City, Panama",
    [
        "Designed relational databases, stored procedures, and integration components for operational applications and reporting workloads.",
        "Delivered end-to-end SQL Server, SSIS, reporting, and data solutions aligned with business and infrastructure requirements.",
    ],
    "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel",
)

role(
    doc,
    "Data Warehouse DBA",
    "BAC Credomatic",
    "November 2017 - January 2019",
    "San Jose, Costa Rica",
    [
        "Developed and maintained SSIS pipelines that extracted, transformed, and loaded multiple source systems into the enterprise data warehouse.",
        "Designed and optimized tables, indexes, and views for reliable, high-performance analytical workloads.",
    ],
    "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure",
)

section_heading(doc, "Additional Database and Integration Experience")
bullet(doc, "Bosal, DBA and BI Consultant: designed the first stage of the enterprise data warehouse and supported SQL Server, SSIS, Azure, MySQL, and Power BI delivery.")
bullet(doc, "Xetux Solutions, Database Manager: established database policies and created a centralized sales data lake for reporting and analysis.")
bullet(doc, "ACH Cloud Services and EducaTablet: administered SQL Server environments, developed data and reporting processes, documented operations, and trained development teams.")
bullet(doc, "VIGEOSOFT and Optica Caroni: modeled OLTP databases, developed SQL and application processes, and delivered operational reporting solutions.")

section_heading(doc, "Education and Languages")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
write(p, "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000 - 2004).")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
write(p, "Additional training: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI; Python for Data Analysis.")
p = doc.add_paragraph()
write(p, "Spanish: Native | English: Full Professional")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "SSIS Integration Engineer CV"
doc.core_properties.subject = "SQL Server, SSIS, ETL, enterprise data integration, testing, and production support"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)

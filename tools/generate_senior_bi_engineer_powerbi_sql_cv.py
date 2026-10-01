from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import generate_fullstack_principal_sql_dba_cv as layout


SOURCE = layout.ROOT / "Base" / "Pedro_Gutierrez_CV.docx"
OUTPUT = layout.ROOT / "Output" / "Pedro_Gutierrez_Senior_BI_Engineer_PowerBI_SQL_CV.docx"


def add_skills_table(doc):
    rows = [
        (
            "Business intelligence",
            "Power BI dashboards and reports, executive and management reporting, KPI reporting, SSRS, SSAS, Excel, requirements analysis and stakeholder communication",
        ),
        (
            "SQL and performance",
            "Advanced T-SQL, complex queries, stored procedures, views, functions, indexing, execution analysis, root-cause investigation and query optimization",
        ),
        (
            "Data modeling",
            "Relational and dimensional modeling, Kimball methodology, fact and dimension design, Data Warehouse architecture, semantic models and reusable reporting datasets",
        ),
        (
            "Python and automation",
            "Python, Pandas, NumPy, Matplotlib, exploratory analysis, validation automation, anomaly detection, recurring workflow automation and PowerShell",
        ),
        (
            "Platforms and integration",
            "SQL Server, Snowflake, dbt, MetricFlow, SSIS, ETL/ELT, Azure-based environments, PostgreSQL, MySQL, Salesforce integrations and Git",
        ),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    widths = [Inches(1.72), Inches(5.68)]
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for index, text in enumerate(("Area", "Tools and practices")):
        cell = table.rows[0].cells[index]
        cell.text = ""
        layout.set_cell_shading(cell, layout.BLUE)
        layout.set_cell_borders(cell)
        layout.set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        layout.write(p, text, 8.7, True, color="FFFFFF")

    for row_index, (area, details) in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate((area, details)):
            cell = cells[index]
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            layout.set_cell_borders(cell)
            layout.set_cell_margins(cell)
            if row_index % 2 == 0:
                layout.set_cell_shading(cell, "F3F7FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            layout.write(p, value, 8.35, bold=(index == 0))


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)
    layout.clear_body(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.43)
    section.bottom_margin = Inches(0.43)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.1)
    normal.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0.5)
    layout.write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 18, True, color=layout.BLUE, font="Aptos Display")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    layout.write(p, "Senior Business Intelligence Engineer | Power BI, SQL and Python", 9.8, True, color="404040")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    layout.write(p, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.5, color="4C4C4C")

    layout.section_heading(doc, "Professional Summary")
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.04
    p.paragraph_format.space_after = Pt(4)
    layout.write(
        p,
        "Senior Business Intelligence and data professional with 15+ years of experience and more than seven years delivering enterprise reporting, analytics, and data warehouse solutions. Strong hands-on background in Power BI, advanced SQL, Python automation, dimensional modeling, ETL/ELT, data quality, and performance optimization. Experienced translating business and operational requirements into reliable reporting datasets, dashboards, KPIs, semantic models, and documented analytical solutions. Proven results include reducing query execution time by up to 50%, reducing manual processing by up to 40%, and improving data-validation accuracy by 30% while collaborating with technical teams, business users, and senior stakeholders.",
    )

    layout.section_heading(doc, "Technical Skills")
    add_skills_table(doc)

    layout.section_heading(doc, "Professional Experience")
    layout.role(
        doc,
        "Snowflake Developer / Data Engineer",
        "ServiceTitan",
        "September 2024 - Present",
        "Remote / Costa Rica",
        [
            "Migrate C# reporting logic into Snowflake SQL, creating maintainable transformations and reporting datasets for analytics workflows.",
            "Develop and maintain dbt models across silver and gold layers, improving data organization, reuse, and downstream reporting reliability.",
            "Optimized Snowflake SQL and dbt workloads, reducing data warehouse processing time by 20% during an initial improvement phase.",
            "Created a Kimball-based data warehouse modeling proof of concept and support MetricFlow semantic models to maintain consistent analytical definitions.",
        ],
        "Snowflake, Snowflake SQL, dbt, MetricFlow, Kimball, dimensional modeling, Data Warehouse, SQL and Git",
    )

    layout.role(
        doc,
        "Data Engineer",
        "Health Catalyst",
        "May 2021 - September 2024",
        "San Jose, Costa Rica",
        [
            "Prepared, cleaned, transformed, and modeled large datasets for Power BI dashboards, business reporting, and operational decision-making.",
            "Designed customized ETL applications that reduced manual data processing time by up to 40%.",
            "Developed Python-based exploratory analysis and validation tools with Pandas, NumPy, and Matplotlib, improving data-validation accuracy by 30%.",
            "Built and supported reporting workflows across SQL Server, SSIS, Snowflake, Power BI, Excel, Python, and PowerShell.",
        ],
        "Power BI, SQL Server, T-SQL, SSIS, Snowflake, Python, Pandas, NumPy, Matplotlib, PowerShell and Excel",
    )

    layout.role(
        doc,
        "SQL Developer",
        "Intertec International",
        "February 2019 - April 2021",
        "San Jose, Costa Rica",
        [
            "Developed and optimized complex SQL queries, stored procedures, and database workflows supporting business logic, reporting, and analytics.",
            "Consolidated SQL Server, Salesforce, and MySQL sources into analytics-ready datasets through optimized ETL processes.",
            "Improved data accuracy by more than 25% and reduced query execution time by up to 50% through validation, error handling, indexing, and SQL tuning.",
        ],
        "SQL Server, advanced T-SQL, SSIS, Salesforce and MySQL",
    )

    layout.role(
        doc,
        "DBA / SQL and BI Consultant",
        "Gold Data Networks",
        "January 2016 - February 2020",
        "Panama City, Panama",
        [
            "Designed relational databases, reporting structures, stored procedures, and integration components for operational and BI requirements.",
            "Delivered SQL Server, SSIS, SSRS, Power BI, PostgreSQL, and Excel solutions aligned with stakeholder and infrastructure needs.",
        ],
        "Power BI, SQL Server, T-SQL, SSIS, SSRS, PostgreSQL and Excel",
        page_break_before=True,
    )

    layout.role(
        doc,
        "Data Warehouse DBA",
        "BAC Credomatic",
        "November 2017 - January 2019",
        "San Jose, Costa Rica",
        [
            "Developed and maintained ETL pipelines integrating multiple source systems into the enterprise data warehouse.",
            "Designed and optimized fact-supporting tables, indexes, views, and analytical structures for reliable, high-performance reporting.",
            "Analyzed large datasets to identify trends and patterns supporting operational and business decision-making.",
        ],
        "Power BI, SQL Server, SSIS, SSAS, SSRS, Excel and Azure",
    )

    layout.role(
        doc,
        "DBA and BI Consultant",
        "Bosal",
        "March 2017 - March 2018",
        "Lummen, Belgium",
        [
            "Designed the first stage of the company's primary data warehouse and developed reports and presentations for senior management.",
            "Partnered with development teams to improve database practices and support BI and reporting delivery.",
        ],
        "Power BI, SQL Server, SSIS, Azure, MySQL and executive reporting",
    )

    layout.section_heading(doc, "Additional Relevant Experience")
    layout.bullet(doc, "Xetux Solutions, Database Manager: established database policies and created a centralized sales data lake for reporting and analysis.", 8.7)
    layout.bullet(doc, "ACH Cloud Services, SQL and BI Consultant: created billing-management dashboards and reports and documented the data environment.", 8.7)
    layout.bullet(doc, "EducaTablet and VIGEOSOFT, SQL Server DBA: developed reporting and database solutions, modeled OLTP structures, and trained development teams.", 8.7)
    layout.bullet(doc, "Optica Caroni, Developer Analyst: delivered operational projects and created management reports and dashboards for production areas.", 8.7)

    layout.section_heading(doc, "Education and Languages")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    layout.write(p, "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000 - 2004)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    layout.write(p, "Training: Analyzing and Visualizing Data with Power BI; Python for Data Analysis; Dataiku Core Designer; R for Data Science")
    p = doc.add_paragraph()
    layout.write(p, "Spanish: Native | English: Full Professional")

    doc.core_properties.author = "Pedro Javier Gutierrez Armas"
    doc.core_properties.title = "Senior Business Intelligence Engineer Resume"
    doc.core_properties.subject = "Resume tailored for a Senior Business Intelligence Engineer role focused on Power BI, SQL, Python, and dimensional modeling"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

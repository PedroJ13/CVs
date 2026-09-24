from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_WWT_Database_Migration_Leader_CV.docx"


def write(paragraph, value, size=9.2, bold=False, color="26333F"):
    run = paragraph.add_run(value)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def section(doc, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    write(p, label.upper(), 10.5, True, "174B70")


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.7)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    write(p, value)


def role(doc, title, company, dates, place, points, technologies, break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = break_before
    write(p, f"{title} | {company}", 10, True, "17365D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    write(p, f"{dates} | {place}", 8.7, color="596878")
    for point in points:
        bullet(doc, point)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write(p, "Technologies: ", 8.6, True, "596878")
    write(p, technologies, 8.6, color="596878")


doc = Document()
page = doc.sections[0]
page.top_margin = Inches(0.52)
page.bottom_margin = Inches(0.52)
page.left_margin = Inches(0.65)
page.right_margin = Inches(0.65)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "PEDRO JAVIER GUTIERREZ ARMAS", 16, True, "17365D")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
write(p, "Database Migration Leader | SQL Server, Azure and Data Platforms", 10.1, True, "174B70")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
write(p, "San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.6, color="596878")

section(doc, "Professional Summary")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.04
write(p, "Database administrator, SQL developer, and data engineer with 15+ years of experience delivering enterprise database, migration, integration, and data warehouse solutions. Led the planning and execution of a ten-database SQL Server migration, coordinating risk identification, data validation, integrity, deployment activities, and operational continuity. Deep experience with SQL Server and T-SQL, plus hands-on work with MySQL, PostgreSQL, Snowflake, Azure-based environments, Python, and PowerShell. Proven in query and workload tuning, technical documentation, database policies, stakeholder coordination, and AI-assisted engineering workflows.")

section(doc, "Core Expertise")
for label, value in [
    ("Database migration", "Migration planning and execution, risk identification, data validation, reconciliation support, integrity, deployment coordination, and operational continuity."),
    ("Database platforms", "Microsoft SQL Server, MySQL, PostgreSQL, Snowflake; schemas, tables, views, indexes, stored procedures, and relational modeling."),
    ("Performance", "Advanced T-SQL, query tuning, indexing strategies, workload optimization, troubleshooting, and post-change validation."),
    ("Automation and data engineering", "Python, PowerShell, SSIS, ETL/ELT, dbt, data quality controls, and repeatable processing workflows."),
    ("Leadership and AI", "Technical coordination, risk communication, documentation, database policies, developer training, Cursor, Snowflake Cortex, and AI-assisted PR review."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    write(p, f"{label}: ", 9.1, True)
    write(p, value, 9.1)

section(doc, "Professional Experience")
role(doc, "Snowflake Developer / Data Engineer", "ServiceTitan", "Sep 2024 - Present", "Remote / Costa Rica", [
    "Migrate C# reporting logic into Snowflake SQL and maintain dbt models across silver and gold data layers.",
    "Optimize Snowflake SQL and dbt workloads, reducing data warehouse processing time by 20% during an initial improvement phase.",
    "Created a Kimball-based warehouse modeling proof of concept and support MetricFlow semantic-layer and Snowflake Cortex workflows.",
], "Snowflake, Snowflake SQL, dbt, MetricFlow, Kimball, Snowflake Cortex, Cursor, Git")
role(doc, "Data Engineer", "SMASH Costa Rica", "May 2021 - Sep 2024", "San Jose, Costa Rica", [
    "Designed production ETL applications and data workflows across SQL Server, SSIS, Snowflake, Python, and PowerShell, reducing manual processing by up to 40%.",
    "Automated data validation and anomaly detection, improving validation accuracy by 30% and strengthening production data quality.",
], "SQL Server, SSIS, Snowflake, Python, PowerShell, Power BI, Pandas")
role(doc, "SQL Developer", "Intertec International", "Feb 2019 - Apr 2021", "San Jose, Costa Rica", [
    "Developed and optimized advanced SQL queries, stored procedures, and data workflows across SQL Server, Salesforce, and MySQL.",
    "Reduced query execution times by up to 50% and improved data accuracy by more than 25% through tuning, indexing, validation, and error handling.",
], "SQL Server, T-SQL, SSIS, Salesforce, MySQL")
role(doc, "DBA / SQL Developer", "EL Tiempo", "Sep 2020 - Nov 2020", "Colombia", [
    "Led planning and execution of migration activities for ten SQL Server databases, delivering all phases on time and within scope.",
    "Coordinated with cross-functional teams to identify migration risks, validate data, preserve integrity, and maintain operational continuity.",
    "Supported multi-platform migration and reporting components throughout deployment activities.",
], "SQL Server, SSIS, Azure, SSRS", break_before=True)
role(doc, "DBA / SQL and BI Consultant", "Gold Data Networks", "Jan 2016 - Feb 2020", "Panama City, Panama", [
    "Designed relational databases, schemas, tables, stored procedures, and supporting objects for operational and reporting workloads.",
    "Delivered end-to-end database, integration, and BI solutions aligned with business and infrastructure requirements.",
], "SQL Server, SSIS, SSRS, PostgreSQL, Power BI, Excel")
role(doc, "Data Warehouse DBA", "BAC Credomatic", "Nov 2017 - Jan 2019", "San Jose, Costa Rica", [
    "Developed and maintained ETL pipelines that loaded multiple source systems into the enterprise data warehouse.",
    "Designed and optimized tables, indexes, and views for reliable, high-performance analytical workloads.",
], "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure")

section(doc, "Additional Leadership and Database Experience")
bullet(doc, "Xetux Solutions, Database Manager: defined database management policies, improved performance and data security, and created a centralized sales data lake.")
bullet(doc, "EducaTablet, SQL Server DBA: administered enterprise databases, improved response times, and trained development teams in database programming practices.")
bullet(doc, "VIGEOSOFT, SQL Server DBA: modeled, designed, configured, and programmed OLTP databases and documented structural changes.")
bullet(doc, "Bosal and ACH Cloud Services: designed warehouse components, administered databases, created management reports, and documented data environments.")

section(doc, "Education and Languages")
p = doc.add_paragraph()
write(p, "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio (2000-2004).")
p = doc.add_paragraph()
write(p, "Additional training: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI; Python for Data Analysis.")
p = doc.add_paragraph()
write(p, "Spanish: Native | English: Full professional proficiency.")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "WWT Database Migration Leader CV"
doc.save(OUTPUT)
print(OUTPUT)

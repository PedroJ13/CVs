from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_Delphos_DBA_SQL_Server_CV.docx"


def text(paragraph, value, size=9.2, bold=False, color="25313D"):
    part = paragraph.add_run(value)
    part.font.name = "Calibri"
    part.font.size = Pt(size)
    part.font.bold = bold
    part.font.color.rgb = RGBColor.from_string(color)


def heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    text(p, title.upper(), 10.5, True, "174B70")


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.6)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    text(p, value)


def role(doc, title, company, dates, location, points, technologies, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4.5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    text(p, f"{title} | {company}", 10, True, "17365D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    text(p, f"{dates} | {location}", 8.7, color="596878")
    for point in points:
        bullet(doc, point)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    text(p, "Tecnologías: ", 8.6, True, "596878")
    text(p, technologies, 8.6, color="596878")


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.52)
sec.bottom_margin = Inches(0.52)
sec.left_margin = Inches(0.65)
sec.right_margin = Inches(0.65)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
text(p, "PEDRO JAVIER GUTIERREZ ARMAS", 16, True, "17365D")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
text(p, "Administrador de Bases de Datos | SQL Server y desarrollo SQL", 10.1, True, "174B70")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
text(p, "San José, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13", 8.6, color="596878")

heading(doc, "Perfil profesional")
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.04
text(p, "Profesional de bases de datos con más de 15 años de experiencia en administración y desarrollo sobre SQL Server en entornos empresariales. Experiencia en diseño de estructuras relacionales, consultas SQL y T-SQL, joins, índices, procedimientos almacenados, ETL, rendimiento, migraciones, integridad de datos y documentación técnica. Ha colaborado con equipos de desarrollo y negocio para resolver problemas de datos y sostener procesos de operación y reportes. Reside en Costa Rica y cuenta con experiencia de trabajo remoto.")

heading(doc, "Competencias técnicas")
for label, value in [
    ("Bases de datos", "SQL Server, PostgreSQL, MySQL, Snowflake; diseño de tablas, vistas e índices."),
    ("SQL y desarrollo", "T-SQL, consultas complejas, joins, procedimientos almacenados, optimización de consultas."),
    ("Administración", "Migraciones, validación e integridad de datos, análisis de rendimiento, documentación de procesos y diccionarios de datos."),
    ("Integración", "SSIS, SSRS, SSAS, ETL, Power BI, Python y PowerShell."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.7)
    text(p, f"{label}: ", 9.1, True)
    text(p, value, 9.1)

heading(doc, "Experiencia profesional")
role(doc, "Snowflake Developer / Data Engineer", "ServiceTitan", "Sep 2024 - Actualidad", "Remoto / Costa Rica", [
    "Migra lógica de reportes desde C# hacia Snowflake SQL y mantiene modelos dbt para datos analíticos reutilizables.",
    "Optimiza consultas y procesamiento del data warehouse; logró una reducción inicial de 20% en el tiempo de procesamiento.",
], "Snowflake SQL, dbt, modelado de datos, Kimball")
role(doc, "Data Engineer", "SMASH Costa Rica", "May 2021 - Sep 2024", "San José, Costa Rica", [
    "Diseñó y mantuvo aplicaciones ETL y flujos de datos con SQL Server y SSIS según requerimientos de clientes.",
    "Automatizó validaciones y análisis de datos, mejorando la precisión de validación en 30% y reduciendo procesamiento manual hasta en 40%.",
], "SQL Server, SSIS, Python, PowerShell, Snowflake, Power BI")
role(doc, "SQL Developer", "Intertec International", "Feb 2019 - Abr 2021", "San José, Costa Rica", [
    "Desarrolló y optimizó procedimientos almacenados, consultas SQL y flujos de bases de datos para lógica empresarial y reportes.",
    "Redujo tiempos de consulta hasta en 50% mediante análisis de rendimiento, optimización SQL y estrategias de índices.",
], "SQL Server, T-SQL, SSIS, Salesforce, MySQL")
role(doc, "DBA / SQL Developer", "EL Tiempo", "Sep 2020 - Nov 2020", "Colombia", [
    "Lideró la migración de diez bases de datos, validando información, riesgos e integridad durante la implementación.",
], "SQL Server, SSIS, Azure, SSRS", page_break_before=True)
role(doc, "DBA / SQL & BI Consultant", "Gold Data Networks", "Ene 2016 - Feb 2020", "Ciudad de Panamá, Panamá", [
    "Construyó bases relacionales, tablas y objetos SQL para aplicaciones y cargas de reportes; desarrolló procedimientos almacenados.",
], "SQL Server, SSIS, SSRS, PostgreSQL, Power BI")
role(doc, "Data Warehouse DBA", "BAC Credomatic", "Nov 2017 - Ene 2019", "San José, Costa Rica", [
    "Mantuvo procesos ETL y definió tablas, índices y vistas para mejorar el rendimiento de cargas analíticas.",
], "SQL Server, SSIS, SSAS, SSRS, Power BI, Azure")

heading(doc, "Experiencia adicional")
bullet(doc, "EducaTablet: administró y desarrolló bases SQL Server y capacitó a desarrolladores en buenas prácticas de programación de datos.")
bullet(doc, "VIGEOSOFT: modeló, diseñó y configuró bases OLTP; documentó cambios estructurales mediante diccionarios de datos.")
bullet(doc, "ACH Cloud Services, Bosal, Xetux Solutions y Optica Caroni: administración de bases de datos, integración y reportes empresariales.")

heading(doc, "Formación e idiomas")
p = doc.add_paragraph()
text(p, "Analista de Sistemas, Informática - IUT Dr. Federico Rivero Palacio (2000-2004).")
p = doc.add_paragraph()
text(p, "Formación adicional: SQL Admin Part 1; Analyzing and Visualizing Data with Power BI.")
p = doc.add_paragraph()
text(p, "Español: nativo | Inglés: dominio profesional.")

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "CV DBA SQL Server Delphos"
doc.save(OUTPUT)
print(OUTPUT)

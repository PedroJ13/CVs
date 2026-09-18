from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Work\CVs\Output\Pedro_Gutierrez_Data_Engineer_BI_Microsoft_Platform_CV.docx"
NAVY = "17365D"
BLUE = "2F75B5"
TEXT = "202938"
MUTED = "52606D"


def font(run, size=9.3, color=TEXT, bold=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def no_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def add_section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    font(r, size=10.4, color=BLUE, bold=True)


def add_bullet(doc, text, size=9.15):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.6)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    font(r, size=size)


def add_role(doc, title, company, dates, location, bullets, technologies, break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5.5)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = break_before
    r = p.add_run(f"{title} | {company}")
    font(r, size=10.1, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{dates} | {location}")
    font(r, size=8.8, color=MUTED)

    for bullet in bullets:
        add_bullet(doc, bullet)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0.5)
    p.paragraph_format.space_after = Pt(1.5)
    r = p.add_run("Technologies: ")
    font(r, size=8.65, color=MUTED, bold=True)
    r = p.add_run(technologies)
    font(r, size=8.65, color=MUTED)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.48)
section.bottom_margin = Inches(0.48)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9.3)
normal.paragraph_format.space_after = Pt(2)

bullet_style = styles["List Bullet"]
bullet_style.font.name = "Aptos"
bullet_style.font.size = Pt(9.15)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(0)
title.style.font.name = "Aptos Display"
r = title.add_run("PEDRO JAVIER GUTIERREZ ARMAS")
font(r, size=16.5, color=NAVY, bold=True, name="Aptos Display")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
r = p.add_run("DATA ENGINEER | BI DEVELOPER | MICROSOFT DATA PLATFORM")
font(r, size=10.1, color=BLUE, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
r = p.add_run("San Jose, Costa Rica | +506 6351-5860 | pj13eros@hotmail.com | linkedin.com/in/pedrogutierrez13")
font(r, size=8.6, color=MUTED)

add_section(doc, "Professional Summary")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
p.paragraph_format.line_spacing = 1.02
summary = (
    "Data Engineer and BI Developer with 15+ years of experience delivering SQL Server, ETL, data warehouse, "
    "reporting, and analytics solutions in consulting and production environments. Advanced background in T-SQL, "
    "SSIS, stored procedures, query tuning, dimensional data structures, Power BI, SSRS, SSAS, Python, PowerShell, "
    "Snowflake, and Azure-based environments. Experienced modernizing reporting logic, integrating heterogeneous "
    "sources, building reliable analytical datasets, supporting database migrations, and translating stakeholder "
    "requirements into maintainable data solutions. Proven results include reducing warehouse processing time by "
    "20%, query execution time by up to 50%, and manual processing by up to 40%."
)
r = p.add_run(summary)
font(r, size=9.2)

add_section(doc, "Technical Skills")
skills = [
    ("Microsoft Data Stack", "SQL Server, T-SQL, SSIS, SSAS, SSRS, Power BI, Excel, Azure-based data environments"),
    ("Data Engineering", "ETL/ELT pipelines, data integration, transformation, orchestration support, automated refresh workflows, production troubleshooting"),
    ("Data Warehousing", "Dimensional modeling, Kimball practices, fact and dimension structures, analytical datasets, data marts, silver and gold layers"),
    ("Development", "Stored procedures, functions, views, complex queries, Python, PowerShell, C# logic migration, Git-based workflows"),
    ("Cloud and Platforms", "Snowflake SQL, dbt models, Snowflake Cortex, MetricFlow semantic layer, PostgreSQL, MySQL, Salesforce integration"),
    ("BI and Quality", "Power BI dashboards, reporting layers, data validation, anomaly detection, error handling, documentation, stakeholder collaboration"),
]
table = doc.add_table(rows=len(skills), cols=2)
table.autofit = False
remove_table_borders(table)
for row, (label, value) in zip(table.rows, skills):
    row.cells[0].width = Inches(1.45)
    row.cells[1].width = Inches(5.75)
    for cell in row.cells:
        no_cell_margins(cell)
    left = row.cells[0].paragraphs[0]
    right = row.cells[1].paragraphs[0]
    left.paragraph_format.space_after = Pt(1)
    right.paragraph_format.space_after = Pt(1)
    r = left.add_run(label)
    font(r, size=8.75, color=NAVY, bold=True)
    r = right.add_run(value)
    font(r, size=8.75)

add_section(doc, "Professional Experience")
add_role(
    doc,
    "Snowflake Developer / Data Engineer",
    "ServiceTitan",
    "Sep 2024 - Present",
    "Remote / Costa Rica",
    [
        "Migrate C# reporting and business logic into Snowflake SQL, creating maintainable transformation and reporting workflows.",
        "Build and maintain dbt models across silver and gold layers to deliver reusable, governed datasets for downstream analytics.",
        "Optimize Snowflake SQL and dbt processing, reducing data warehouse processing time by approximately 20% in the initial optimization phase.",
        "Created a Kimball-based data warehouse modeling proof of concept and support MetricFlow semantic-layer development with Snowflake Cortex.",
    ],
    "Snowflake, Snowflake SQL, dbt, MetricFlow, Snowflake Cortex, dimensional modeling, SQL, C# logic migration, Git",
)

add_role(
    doc,
    "Data Engineer",
    "SMASH Costa Rica",
    "May 2021 - Sep 2024",
    "San Jose, Costa Rica",
    [
        "Designed client-specific ETL applications and automated data workflows, reducing manual processing time by up to 40%.",
        "Built and maintained SQL Server and SSIS integration processes using Python, PowerShell, Snowflake, Power BI, and Excel across diverse source systems.",
        "Prepared, transformed, and modeled reporting datasets for Power BI dashboards and operational decision-making.",
        "Automated validation and exploratory analysis with Pandas, NumPy, and Matplotlib, improving data validation accuracy by 30%.",
    ],
    "SQL Server, T-SQL, SSIS, Power BI, Python, PowerShell, Snowflake, Pandas, NumPy, Excel",
)

add_role(
    doc,
    "SQL Developer",
    "Intertec International",
    "Feb 2019 - Apr 2021",
    "San Jose, Costa Rica",
    [
        "Developed complex T-SQL queries, stored procedures, and database workflows supporting business logic, reporting, and data integration.",
        "Consolidated SQL Server, Salesforce, and MySQL data into analytics-ready datasets through optimized ETL processes.",
        "Reduced query execution times by up to 50% through execution analysis, indexing strategies, and SQL performance tuning.",
        "Improved data accuracy by more than 25% through validation controls and error-handling mechanisms.",
    ],
    "SQL Server, T-SQL, SSIS, Salesforce, MySQL, stored procedures, query optimization",
)

add_role(
    doc,
    "DBA / SQL Developer",
    "EL Tiempo",
    "Sep 2020 - Nov 2020",
    "Colombia",
    [
        "Led the planning and execution of migration activities for 10 databases while preserving data integrity and operational continuity.",
        "Partnered with technical stakeholders to identify risks, validate data, and support deployment and reporting activities.",
    ],
    "SQL Server, SSIS, Azure, SSRS, database migration",
    break_before=True,
)

add_role(
    doc,
    "DBA / SQL and BI Consultant",
    "Gold Data Networks",
    "Jan 2016 - Feb 2020",
    "Panama City, Panama",
    [
        "Designed relational databases, table structures, stored procedures, and end-to-end data solutions for application and BI workloads.",
        "Delivered SQL Server, SSIS, SSRS, PostgreSQL, Power BI, and Excel solutions integrated with client infrastructure and reporting needs.",
        "Worked directly with stakeholders to translate requirements into scalable database, integration, and reporting solutions.",
    ],
    "SQL Server, T-SQL, SSIS, SSRS, PostgreSQL, Power BI, Excel",
)

add_role(
    doc,
    "Data Warehouse DBA",
    "BAC Credomatic",
    "Nov 2017 - Jan 2019",
    "San Jose, Costa Rica",
    [
        "Developed and maintained SSIS pipelines that loaded data from multiple sources into the enterprise data warehouse.",
        "Designed and optimized tables, indexes, and views for analytical workloads and supported reporting through SSAS, Power BI, and SSRS.",
        "Analyzed large datasets to identify trends and support data-driven operational decisions.",
    ],
    "SQL Server, T-SQL, SSIS, SSAS, Power BI, SSRS, Azure, data warehousing",
)

add_section(doc, "Additional Relevant Experience")
add_bullet(doc, "Designed the first stage of Bosal's main data warehouse and produced management reporting for senior stakeholders.")
add_bullet(doc, "Created a centralized sales data lake at Xetux Solutions and supported SQL Server, SSIS, SSRS, and Azure-based reporting environments across multiple consulting roles.")
add_bullet(doc, "Developed databases, ETL processes, operational reports, dashboards, and technical documentation at ACH Cloud Services, EducaTablet, VIGEOSOFT, and Optica Caroni.")

add_section(doc, "Education and Training")
add_bullet(doc, "Systems Analyst, Informatics - IUT Dr. Federico Rivero Palacio")
add_bullet(doc, "English Certificate - Universidad Central de Venezuela / Microsoft")
add_bullet(doc, "Analyzing and Visualizing Data with Power BI | SQL Admin Part 1 | Dataiku Core Designer")
add_bullet(doc, "Complete Data Science Training with Python for Data Analysis | R Programming A-Z")

add_section(doc, "Languages")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Spanish: Native | English: Full professional proficiency")
font(r, size=9.2)

doc.core_properties.author = "Pedro Javier Gutierrez Armas"
doc.core_properties.title = "Pedro Gutierrez Data Engineer BI Developer Microsoft Data Platform CV"
doc.core_properties.subject = "Data Engineer and BI Developer resume"
doc.save(OUTPUT)
print(OUTPUT)

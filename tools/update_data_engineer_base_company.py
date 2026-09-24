from pathlib import Path

from docx import Document


ROOT = Path(r"C:\Work\CVs")
PATH = ROOT / "Base" / "Pedro_Gutierrez_CV.docx"
OLD = "SMASH Costa Rica"
NEW = "Health Catalyst"


def replace_in_paragraph(paragraph):
    for run in paragraph.runs:
        if OLD in run.text:
            run.text = run.text.replace(OLD, NEW)


def replace_in_table(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(paragraph)
            for nested in cell.tables:
                replace_in_table(nested)


doc = Document(PATH)
for paragraph in doc.paragraphs:
    replace_in_paragraph(paragraph)
for table in doc.tables:
    replace_in_table(table)
for section in doc.sections:
    for part in (section.header, section.footer):
        for paragraph in part.paragraphs:
            replace_in_paragraph(paragraph)
        for table in part.tables:
            replace_in_table(table)

doc.save(PATH)
print(PATH)
